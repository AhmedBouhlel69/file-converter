"""Capture deterministic offscreen screenshots of the redesigned UI states.

Run from the repository root:
    python -m image_converter.ui.dev_capture_states
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

# This must be configured before any Qt import.
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from PySide6.QtCore import QThreadPool, Qt
from PySide6.QtGui import QFontDatabase
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QDialog, QGridLayout, QLabel, QSpinBox,
)

from image_converter.ui.main_window import ImageConverterMainWindow
from image_converter.ui.theme import apply_theme, set_variant
from image_converter.ui.tool_dialogs import (
    CompressDialog, ImageCompressDialog, MergeDialog, OrganizePagesDialog, SplitDialog,
)


def _load_offscreen_font() -> None:
    """The Windows offscreen platform does not always discover system fonts."""
    if sys.platform == "win32":
        segoe = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "segoeui.ttf"
        if segoe.exists():
            QFontDatabase.addApplicationFont(str(segoe))


def _create_samples(folder: Path) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)

    image_path = folder / "campaign-hero.png"
    image = Image.new("RGB", (1280, 720), "#6654E8")
    draw = ImageDraw.Draw(image)
    draw.rectangle((90, 90, 1190, 630), outline="#67DFFF", width=12)
    image.save(image_path)

    pdf_path = folder / "quarterly-report.pdf"
    pdf = fitz.open()
    for number in range(1, 4):
        page = pdf.new_page()
        page.insert_text((72, 96), f"Quarterly report — page {number}", fontsize=20)
    pdf_path.write_bytes(pdf.tobytes())
    pdf.close()

    document_path = folder / "meeting-notes.docx"
    document = Document()
    document.add_heading("Meeting notes", level=1)
    document.add_paragraph("A representative document used for UI state capture.")
    document.save(document_path)

    sheet_path = folder / "regional-sales.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Region", "Revenue"])
    sheet.append(["North", 42000])
    sheet.append(["South", 37000])
    workbook.save(sheet_path)

    return [image_path, pdf_path, document_path, sheet_path]


def _settle(app: QApplication, milliseconds: int = 250) -> None:
    deadline = time.monotonic() + milliseconds / 1000
    while time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    app.processEvents()


def capture_states(output_dir: Path, width: int = 1380, height: int = 860) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    app = QApplication.instance() or QApplication([])
    _load_offscreen_font()
    apply_theme(app)

    window = ImageConverterMainWindow()
    window.resize(width, height)
    window.show()
    captured: list[Path] = []

    def capture(name: str) -> None:
        _settle(app)
        window.repaint()
        app.processEvents()
        destination = output_dir / f"{name}.png"
        if not window.grab().save(str(destination)):
            raise RuntimeError(f"Could not save {destination}")
        captured.append(destination)

    capture("01-empty")

    window.drag_overlay.setGeometry(window.centralWidget().rect().adjusted(4, 4, -4, -4))
    window.drag_overlay.show()
    window.drag_overlay.raise_()
    capture("02-drag-over")
    window.drag_overlay.hide()

    samples = _create_samples(output_dir / "samples")
    window.add_images(samples[:1])
    _settle(app, 500)
    capture("03-single-type")
    window.add_images(samples[1:])
    _settle(app, 1100)
    capture("03-populated-mixed")

    window.queue_table.selectRow(0)
    _settle(app)
    capture("04-row-selected")
    window._show_batch_settings()

    model = window.queue_table.queue_model
    window._is_converting = True
    window._footer_state = "converting"
    window._active_rows = list(range(model.rowCount()))
    for row in range(model.rowCount()):
        model.update_status(row, "Converting" if row == 0 else "Queued", progress=36 if row == 0 else 0)
    window._set_queue_editing_enabled(False)
    window.progress_bar.setValue(36)
    window.progress_bar.show()
    window.lbl_status.setText(f"Converting 1 of {model.rowCount()} files · 36%")
    window.lbl_status_hint.setText("Cancel stops after the current file finishes.")
    window._refresh_queue_ui(update_status=False)
    capture("05-converting")

    window._is_converting = False
    window._set_queue_editing_enabled(True)
    for row, target in enumerate(("WEBP", "TXT", "TXT", "CSV")):
        window.queue_table.set_target_format_for_row(row, target)
    window.start_conversion()
    deadline = time.monotonic() + 20
    while window._is_converting and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    if window._is_converting:
        window._cancel_conversion()
        raise RuntimeError("Sample conversion did not finish within 20 seconds")
    if any(entry.status not in ("Done", "Warning") for entry in model.entries):
        raise RuntimeError("A sample conversion failed; completed captures must use real outputs")
    capture("06-completed")

    model.update_status(0, "Failed", progress=0, error="The sample file could not be decoded.")
    window._footer_state = "partial_failure"
    window.btn_open_folder.setVisible(model.rowCount() > 1)
    window.lbl_status.setText(f"{model.rowCount() - 1} converted · 1 failed")
    window.lbl_status_hint.setText("Select a failed status to retry.")
    set_variant(window.btn_convert_all, "primary")
    window._refresh_queue_ui(update_status=False)
    capture("07-error")

    window._set_app_mode(1)
    capture("08-image-tools")
    window._set_app_mode(2)
    for index, name in enumerate(("merge", "split", "organize", "compress")):
        window.pdf_tools_page.set_active_tool(index)
        capture(f"09-pdf-{name}")
    window._set_app_mode(3)
    window.background_removal_page.add_files([samples[0]])
    capture("10-background")
    window.background_removal_page.btn_advanced.setChecked(True)
    capture("11-background-advanced")

    window.close()
    QThreadPool.globalInstance().waitForDone(3000)
    app.processEvents()
    captured.extend(capture_dialogs(output_dir, samples, width, height))
    (output_dir / "manifest.json").write_text(json.dumps({
        "logical_size": [width, height],
        "platform": app.platformName(),
        "captures": [path.name for path in captured],
        "fixtures": "Running and failure queue states are layout fixtures. Completed sizes come from a real sample batch.",
        "native_windows_dpi_or_assistive_technology_verified": False,
    }, indent=2), encoding="utf-8")
    return captured


def capture_dialogs(output_dir: Path, samples: list[Path], width: int, height: int) -> list[Path]:
    """Render every utility dialog plus checked/mixed/disabled/focused controls."""
    app = QApplication.instance()
    captured = []
    dialogs = (
        ("merge", lambda: MergeDialog(initial_files=samples[:2])),
        ("split", lambda: SplitDialog(initial_file=samples[1])),
        ("organize", lambda: OrganizePagesDialog(initial_file=samples[1])),
        ("compress-pdf", lambda: CompressDialog(initial_file=samples[1])),
        ("compress-images", lambda: ImageCompressDialog(initial_files=[samples[0]])),
    )
    for name, factory in dialogs:
        dialog = factory()
        dialog.resize(width, height)
        dialog.show()
        _settle(app)
        path = output_dir / f"12-dialog-{name}.png"
        if not dialog.grab().save(str(path)):
            raise RuntimeError(f"Could not save {path}")
        captured.append(path)
        dialog.close()
    dialog = QDialog()
    dialog.setWindowTitle("Control state audit")
    apply_theme(dialog)
    grid = QGridLayout(dialog)
    for row, (name, state, enabled) in enumerate((
        ("Unchecked", Qt.CheckState.Unchecked, True),
        ("Checked", Qt.CheckState.Checked, True),
        ("Indeterminate", Qt.CheckState.PartiallyChecked, True),
        ("Disabled unchecked", Qt.CheckState.Unchecked, False),
        ("Disabled checked", Qt.CheckState.Checked, False),
        ("Disabled indeterminate", Qt.CheckState.PartiallyChecked, False),
    )):
        check = QCheckBox(name)
        check.setAccessibleName(name)
        check.setTristate(True)
        check.setCheckState(state)
        check.setEnabled(enabled)
        grid.addWidget(check, row, 0)
        spin = QSpinBox()
        spin.setAccessibleName(f"{name} stepper")
        spin.setRange(0, 100)
        spin.setValue(42)
        spin.setEnabled(enabled)
        grid.addWidget(spin, row, 1)
    focused = QCheckBox("Keyboard focus")
    focused.setAccessibleName("Focused checkbox")
    focused.setChecked(True)
    grid.addWidget(focused, 6, 0)
    for row, (name, value) in enumerate((("Minimum", 0), ("Maximum", 100)), start=7):
        grid.addWidget(QLabel(name), row, 0)
        spin = QSpinBox()
        spin.setAccessibleName(f"{name} stepper")
        spin.setRange(0, 100)
        spin.setValue(value)
        grid.addWidget(spin, row, 1)
    grid.addWidget(QLabel("Tab to inspect checkbox and spinbox focus."), 9, 0, 1, 2)
    dialog.resize(width, height)
    dialog.show()
    focused.setFocus(Qt.FocusReason.TabFocusReason)
    _settle(app)
    path = output_dir / "13-control-states.png"
    if not dialog.grab().save(str(path)):
        raise RuntimeError(f"Could not save {path}")
    captured.append(path)
    dialog.close()
    return captured


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("scratch") / "ui-state-captures",
        help="Directory for PNG captures and generated sample files.",
    )
    parser.add_argument("--width", type=int, default=1380, help="Logical capture width.")
    parser.add_argument("--height", type=int, default=860, help="Logical capture height.")
    args = parser.parse_args()
    for path in capture_states(args.output.resolve(), args.width, args.height):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
