"""
Unit tests for Data conversions (CSV, XLSX).
"""

from pathlib import Path
import json
import tempfile
import fitz
import openpyxl
import pandas as pd
import pytest

from image_converter.core.engine import (
    ConversionConfig,
    ImageConverterEngine,
    get_file_metadata,
)


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


@pytest.fixture
def engine():
    return ImageConverterEngine()


def create_sample_csv(path: Path) -> Path:
    """Create a sample CSV file with test data."""
    df = pd.DataFrame({
        "ID": [101, 102, 103],
        "Product": ["Laptop", "Monitor", "Keyboard"],
        "Price": [1200.50, 299.99, 89.00],
        "InStock": [True, True, False],
    })
    df.to_csv(str(path), index=False)
    return path


def create_sample_xlsx(path: Path) -> Path:
    """Create a sample XLSX file with openpyxl."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Units", "Revenue"])
    ws.append(["North", 45, 9000])
    ws.append(["South", 70, 14000])
    ws.append(["East", 30, 6000])
    wb.save(str(path))
    return path


def test_csv_metadata(temp_dir):
    csv_path = temp_dir / "test.csv"
    create_sample_csv(csv_path)

    meta = get_file_metadata(csv_path)
    assert meta["format"] == "CSV"
    assert meta["rows"] == 3
    assert meta["cols"] == 4
    assert "Product" in meta["columns"]


def test_xlsx_metadata(temp_dir):
    xlsx_path = temp_dir / "test.xlsx"
    create_sample_xlsx(xlsx_path)

    meta = get_file_metadata(xlsx_path)
    assert meta["format"] == "XLSX"
    assert "Sales" in meta["sheets"]
    assert meta["rows"] >= 4
    assert meta["cols"] >= 3


def test_csv_to_xlsx(engine, temp_dir):
    csv_path = temp_dir / "input.csv"
    create_sample_csv(csv_path)
    out_xlsx = temp_dir / "output.xlsx"

    config = ConversionConfig(target_format="XLSX")
    result = engine.convert_single(csv_path, out_xlsx, config)

    assert result.success is True
    assert out_xlsx.is_file()

    # Verify content in created Excel workbook
    wb = openpyxl.load_workbook(str(out_xlsx))
    ws = wb.active
    assert ws.cell(1, 1).value == "ID"
    assert ws.cell(2, 2).value == "Laptop"
    wb.close()


def test_csv_to_pdf(engine, temp_dir):
    csv_path = temp_dir / "input.csv"
    create_sample_csv(csv_path)
    out_pdf = temp_dir / "output.pdf"

    config = ConversionConfig(target_format="PDF")
    result = engine.convert_single(csv_path, out_pdf, config)

    assert result.success is True
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 0

    doc = fitz.open(str(out_pdf))
    assert len(doc) >= 1
    page_text = doc[0].get_text()
    assert "Laptop" in page_text or "Monitor" in page_text
    doc.close()


def test_csv_to_json(engine, temp_dir):
    csv_path = temp_dir / "input.csv"
    create_sample_csv(csv_path)
    out_json = temp_dir / "output.json"

    config = ConversionConfig(target_format="JSON")
    result = engine.convert_single(csv_path, out_json, config)

    assert result.success is True
    assert out_json.is_file()

    data = json.loads(out_json.read_text(encoding="utf-8"))
    assert isinstance(data, list)
    assert len(data) == 3
    assert data[0]["Product"] == "Laptop"


def test_csv_to_html(engine, temp_dir):
    csv_path = temp_dir / "input.csv"
    create_sample_csv(csv_path)
    out_html = temp_dir / "output.html"

    config = ConversionConfig(target_format="HTML")
    result = engine.convert_single(csv_path, out_html, config)

    assert result.success is True
    assert out_html.is_file()
    content = out_html.read_text(encoding="utf-8")
    assert "<table" in content
    assert "Laptop" in content


def test_csv_to_txt(engine, temp_dir):
    csv_path = temp_dir / "input.csv"
    create_sample_csv(csv_path)
    out_txt = temp_dir / "output.txt"

    config = ConversionConfig(target_format="TXT")
    result = engine.convert_single(csv_path, out_txt, config)

    assert result.success is True
    assert out_txt.is_file()
    content = out_txt.read_text(encoding="utf-8")
    assert "Laptop" in content


def test_xlsx_to_csv(engine, temp_dir):
    xlsx_path = temp_dir / "input.xlsx"
    create_sample_xlsx(xlsx_path)
    out_csv = temp_dir / "output.csv"

    config = ConversionConfig(target_format="CSV")
    result = engine.convert_single(xlsx_path, out_csv, config)

    assert result.success is True
    assert out_csv.is_file()

    df = pd.read_csv(str(out_csv))
    assert "Region" in df.columns
    assert len(df) == 3
    assert df.iloc[0]["Region"] == "North"


def test_xlsx_to_pdf(engine, temp_dir):
    xlsx_path = temp_dir / "input.xlsx"
    create_sample_xlsx(xlsx_path)
    out_pdf = temp_dir / "output.pdf"

    config = ConversionConfig(target_format="PDF")
    result = engine.convert_single(xlsx_path, out_pdf, config)

    assert result.success is True
    assert out_pdf.is_file()

    doc = fitz.open(str(out_pdf))
    assert len(doc) >= 1
    doc.close()


def test_xlsx_to_json(engine, temp_dir):
    xlsx_path = temp_dir / "input.xlsx"
    create_sample_xlsx(xlsx_path)
    out_json = temp_dir / "output.json"

    config = ConversionConfig(target_format="JSON")
    result = engine.convert_single(xlsx_path, out_json, config)

    assert result.success is True
    assert out_json.is_file()

    data = json.loads(out_json.read_text(encoding="utf-8"))
    assert "Sales" in data or isinstance(data, list)


def test_xlsx_to_html(engine, temp_dir):
    xlsx_path = temp_dir / "input.xlsx"
    create_sample_xlsx(xlsx_path)
    out_html = temp_dir / "output.html"

    config = ConversionConfig(target_format="HTML")
    result = engine.convert_single(xlsx_path, out_html, config)

    assert result.success is True
    assert out_html.is_file()
    content = out_html.read_text(encoding="utf-8")
    assert "<table" in content
    assert "Revenue" in content
