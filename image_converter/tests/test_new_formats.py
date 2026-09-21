"""
Unit tests for new format support: PPTX (PowerPoint), ODT (OpenDocument), and RTF.
"""

from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
import docx
from odf import opendocument, text as odf_text

from image_converter.core.engine import (
    ConversionConfig,
    ImageConverterEngine,
    get_file_metadata,
)
from image_converter.core.presentation_engine import (
    PresentationConverter,
    get_pptx_metadata,
)
from image_converter.core.rich_doc_engine import (
    RichDocumentConverter,
    get_odt_metadata,
    get_rtf_metadata,
)


@pytest.fixture
def engine():
    return ImageConverterEngine()


def create_sample_pptx(path: Path) -> Path:
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]  # Title slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Universal Converter PPTX Test"
    slide.shapes.placeholders[1].text = "Subtitle: Robust Multi-Format Testing"

    # Slide 2: Bullet points
    layout2 = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(layout2)
    slide2.shapes.title.text = "Key Highlights"
    tf = slide2.shapes.placeholders[1].text_frame
    tf.text = "Bullet 1: Pure-Python processing"
    p2 = tf.add_paragraph()
    p2.text = "Bullet 2: Zero Office dependencies"

    prs.save(str(path))
    return path


def create_sample_odt(path: Path) -> Path:
    doc = opendocument.OpenDocumentText()
    p1 = odf_text.P(text="ODT Document Header")
    p2 = odf_text.P(text="Paragraph in OpenDocument format for verification.")
    doc.text.addElement(p1)
    doc.text.addElement(p2)
    doc.save(str(path))
    return path


def create_sample_rtf(path: Path) -> Path:
    rtf_content = (
        r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Calibri;}}"
        r"\f0\fs24 RTF Header Title\par"
        r"This is a paragraph inside a Rich Text Format document.\par}"
    )
    path.write_text(rtf_content, encoding="utf-8")
    return path


# -----------------------------------------------------------------------------
# PPTX Tests
# -----------------------------------------------------------------------------

def test_pptx_metadata_and_preview(tmp_path: Path):
    pptx_path = tmp_path / "presentation.pptx"
    create_sample_pptx(pptx_path)

    meta = get_pptx_metadata(pptx_path)
    assert meta["format"] == "PPTX"
    assert meta["slide_count"] == 2
    assert "Universal Converter" in meta["title"]


@pytest.mark.office
def test_pptx_to_pdf(engine, tmp_path: Path):
    pptx_path = tmp_path / "pres.pptx"
    create_sample_pptx(pptx_path)

    out_pdf = tmp_path / "pres.pdf"
    res = engine.convert_single(pptx_path, out_pdf, ConversionConfig(target_format="PDF"))
    assert res.success is True
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 0


def test_pptx_to_txt(engine, tmp_path: Path):
    pptx_path = tmp_path / "pres.pptx"
    create_sample_pptx(pptx_path)

    out_txt = tmp_path / "pres.txt"
    res = engine.convert_single(pptx_path, out_txt, ConversionConfig(target_format="TXT"))
    assert res.success is True
    assert out_txt.is_file()
    content = out_txt.read_text(encoding="utf-8")
    assert "Universal Converter PPTX Test" in content
    assert "Slide 2" in content


def test_pptx_to_html(engine, tmp_path: Path):
    pptx_path = tmp_path / "pres.pptx"
    create_sample_pptx(pptx_path)

    out_html = tmp_path / "pres.html"
    res = engine.convert_single(pptx_path, out_html, ConversionConfig(target_format="HTML"))
    assert res.success is True
    assert out_html.is_file()
    assert "<html" in out_html.read_text(encoding="utf-8").lower()


def test_text_to_pptx(engine, tmp_path: Path):
    txt_path = tmp_path / "outline.txt"
    txt_path.write_text("=== Slide 1 ===\nIntro Slide\nWelcome to UFC\n\n=== Slide 2 ===\nAgenda\nFeatures", encoding="utf-8")

    out_pptx = tmp_path / "generated.pptx"
    res = engine.convert_single(txt_path, out_pptx, ConversionConfig(target_format="PPTX"))
    assert res.success is True
    assert out_pptx.is_file()

    prs = Presentation(str(out_pptx))
    assert len(prs.slides) == 2


# -----------------------------------------------------------------------------
# ODT Tests
# -----------------------------------------------------------------------------

def test_odt_metadata(tmp_path: Path):
    odt_path = tmp_path / "doc.odt"
    create_sample_odt(odt_path)

    meta = get_odt_metadata(odt_path)
    assert meta["format"] == "ODT"
    assert meta["paragraphs"] == 2
    assert meta["word_count"] > 5


def test_odt_to_pdf_and_txt(engine, tmp_path: Path):
    odt_path = tmp_path / "doc.odt"
    create_sample_odt(odt_path)

    # To PDF
    out_pdf = tmp_path / "doc.pdf"
    res_pdf = engine.convert_single(odt_path, out_pdf, ConversionConfig(target_format="PDF"))
    assert res_pdf.success is True
    assert out_pdf.is_file()

    # To TXT
    out_txt = tmp_path / "doc.txt"
    res_txt = engine.convert_single(odt_path, out_txt, ConversionConfig(target_format="TXT"))
    assert res_txt.success is True
    assert "OpenDocument format" in out_txt.read_text(encoding="utf-8")


@pytest.mark.office
def test_docx_to_odt(engine, tmp_path: Path):
    docx_path = tmp_path / "source.docx"
    d = docx.Document()
    d.add_heading("Docx to ODT", level=1)
    d.add_paragraph("Paragraph inside docx to be converted into ODT.")
    d.save(str(docx_path))

    out_odt = tmp_path / "converted.odt"
    res = engine.convert_single(docx_path, out_odt, ConversionConfig(target_format="ODT"))
    assert res.success is True
    assert out_odt.is_file()


# -----------------------------------------------------------------------------
# RTF Tests
# -----------------------------------------------------------------------------

def test_rtf_metadata(tmp_path: Path):
    rtf_path = tmp_path / "doc.rtf"
    create_sample_rtf(rtf_path)

    meta = get_rtf_metadata(rtf_path)
    assert meta["format"] == "RTF"
    assert meta["word_count"] > 5


@pytest.mark.office
def test_rtf_to_txt_and_docx(engine, tmp_path: Path):
    rtf_path = tmp_path / "doc.rtf"
    create_sample_rtf(rtf_path)

    # To TXT
    out_txt = tmp_path / "doc.txt"
    res_txt = engine.convert_single(rtf_path, out_txt, ConversionConfig(target_format="TXT"))
    assert res_txt.success is True
    assert "RTF Header Title" in out_txt.read_text(encoding="utf-8")

    # To DOCX
    out_docx = tmp_path / "doc.docx"
    res_docx = engine.convert_single(rtf_path, out_docx, ConversionConfig(target_format="DOCX"))
    assert res_docx.success is True
    assert out_docx.is_file()
