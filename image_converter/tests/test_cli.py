"""
Tests for Universal File Converter CLI.
"""

from pathlib import Path
import docx
import fitz
import pandas as pd
from PIL import Image
import pytest

try:
    import pillow_heif
except ImportError:
    pillow_heif = None
else:
    pillow_heif.register_heif_opener()

from image_converter.cli import main, parse_color, collect_image_files, collect_convertible_files


def test_parse_color():
    assert parse_color("#ffffff") == (255, 255, 255)
    assert parse_color("#000") == (0, 0, 0)
    assert parse_color("255, 0, 128") == (255, 0, 128)
    with pytest.raises(ValueError):
        parse_color("invalid")
    with pytest.raises(ValueError):
        parse_color("256,0,0")
    with pytest.raises(ValueError):
        parse_color("1,2,3,4")


def test_cli_rejects_invalid_color():
    with pytest.raises(SystemExit):
        main(["--bg-color", "256,0,0"])


def test_cli_list_formats(capsys):
    exit_code = main(["--list-formats"])
    assert exit_code == 0
    captured = capsys.readouterr().out
    assert "Images" in captured
    assert "Documents" in captured
    assert "Spreadsheets" in captured


def test_cli_batch_flow(tmp_path: Path):
    if pillow_heif is None:
        pytest.skip("HEIF library not available in this test environment")

    in_dir = tmp_path / "inputs"
    out_dir = tmp_path / "outputs"
    in_dir.mkdir()
    out_dir.mkdir()

    # Create dummy images
    Image.new("RGB", (60, 60), color="red").save(in_dir / "img1.jpg")
    Image.new("RGBA", (50, 50), color=(0, 255, 0, 128)).save(in_dir / "img2.png")
    Image.new("RGB", (40, 40), color="blue").save(in_dir / "img3.heic", format="HEIF")

    # Run CLI conversion to PNG
    exit_code = main(["-i", str(in_dir), "-f", "PNG", "-o", str(out_dir)])
    assert exit_code == 0

    outputs = list(out_dir.glob("*.png"))
    assert len(outputs) == 3
    names = {f.name for f in outputs}
    assert "img1.png" in names
    assert "img2.png" in names
    assert "img3.png" in names


def test_cli_csv_to_xlsx(tmp_path: Path):
    csv_file = tmp_path / "data.csv"
    xlsx_file = tmp_path / "data.xlsx"

    pd.DataFrame({"Item": ["A", "B"], "Val": [10, 20]}).to_csv(str(csv_file), index=False)

    exit_code = main(["-i", str(csv_file), "-o", str(xlsx_file)])
    assert exit_code == 0
    assert xlsx_file.is_file()


def test_cli_pdf_to_txt(tmp_path: Path):
    pdf_file = tmp_path / "doc.pdf"
    txt_file = tmp_path / "doc.txt"

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "CLI extracted PDF text", fontsize=12)
    doc.save(str(pdf_file))
    doc.close()

    exit_code = main(["-i", str(pdf_file), "-o", str(txt_file)])
    assert exit_code == 0
    assert txt_file.is_file()
    assert "CLI extracted PDF text" in txt_file.read_text(encoding="utf-8")


def test_cli_docx_to_pdf(tmp_path: Path):
    docx_file = tmp_path / "sample.docx"
    pdf_file = tmp_path / "sample.pdf"

    doc = docx.Document()
    doc.add_heading("CLI Docx Test", level=1)
    doc.add_paragraph("CLI Paragraph test.")
    doc.save(str(docx_file))

    exit_code = main(["-i", str(docx_file), "-o", str(pdf_file)])
    assert exit_code == 0
    assert pdf_file.is_file()
    assert pdf_file.stat().st_size > 0


def test_cli_rejects_duplicate_destinations(tmp_path: Path, capsys):
    in_dir = tmp_path / "inputs"
    out_dir = tmp_path / "outputs"
    in_dir.mkdir()
    out_dir.mkdir()
    Image.new("RGB", (20, 10), color="red").save(in_dir / "same.jpg")
    Image.new("RGB", (30, 15), color="blue").save(in_dir / "same.png")

    exit_code = main(["-i", str(in_dir), "-f", "JPG", "-o", str(out_dir)])

    assert exit_code == 1
    assert "would overwrite" in capsys.readouterr().err
