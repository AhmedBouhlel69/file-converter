"""
Unit tests for Document conversions (PDF, DOCX).
"""

from pathlib import Path
import tempfile
import fitz
import docx
import pytest
from PIL import Image

from image_converter.core.engine import (
    ConversionConfig,
    ImageConverterEngine,
    get_file_metadata,
)
from image_converter.core.document_engine import DocumentConverter


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


@pytest.fixture
def engine():
    return ImageConverterEngine()


def create_sample_pdf(path: Path, text: str = "Hello Universal File Converter") -> Path:
    """Create a simple PDF with test text."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 72), text, fontsize=14)
    doc.save(str(path))
    doc.close()
    return path


def create_sample_docx(path: Path, title: str = "Test Document", body: str = "This is a test paragraph.") -> Path:
    """Create a simple DOCX file with headings, paragraphs, and a table."""
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)

    # Add a 2x2 table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Cell A"
    table.cell(1, 1).text = "Cell B"

    doc.save(str(path))
    return path


def test_pdf_metadata(temp_dir):
    pdf_path = temp_dir / "meta.pdf"
    create_sample_pdf(pdf_path, "Metadata test content")

    meta = get_file_metadata(pdf_path)
    assert meta["format"] == "PDF"
    assert meta["page_count"] == 1
    assert meta["width"] == 595
    assert meta["height"] == 842
    assert meta["file_size"] > 0


def test_docx_metadata(temp_dir):
    docx_path = temp_dir / "meta.docx"
    create_sample_docx(docx_path, "Docx Title", "Some sample body text.")

    meta = get_file_metadata(docx_path)
    assert meta["format"] == "DOCX"
    assert meta["paragraphs"] >= 1
    assert meta["tables"] == 1
    assert meta["word_count"] > 0


def test_pdf_to_png(engine, temp_dir):
    pdf_path = temp_dir / "sample.pdf"
    create_sample_pdf(pdf_path, "PDF to PNG test")
    out_img = temp_dir / "rendered.png"

    config = ConversionConfig(target_format="PNG")
    result = engine.convert_single(pdf_path, out_img, config)

    assert result.success is True
    assert out_img.is_file()
    assert out_img.stat().st_size > 0

    with Image.open(out_img) as img:
        assert img.format == "PNG"
        assert img.width > 0
        assert img.height > 0


def test_pdf_to_txt(engine, temp_dir):
    pdf_path = temp_dir / "sample.pdf"
    create_sample_pdf(pdf_path, "Unique keyword text content")
    out_txt = temp_dir / "sample.txt"

    config = ConversionConfig(target_format="TXT")
    result = engine.convert_single(pdf_path, out_txt, config)

    assert result.success is True
    assert out_txt.is_file()
    content = out_txt.read_text(encoding="utf-8")
    assert "Unique keyword text content" in content


def test_pdf_to_docx(engine, temp_dir):
    pdf_path = temp_dir / "sample.pdf"
    create_sample_pdf(pdf_path, "PDF to Word conversion test")
    out_docx = temp_dir / "sample.docx"

    config = ConversionConfig(target_format="DOCX")
    result = engine.convert_single(pdf_path, out_docx, config)

    assert result.success is True
    assert out_docx.is_file()
    assert out_docx.stat().st_size > 0

    # Verify readable with python-docx
    doc = docx.Document(str(out_docx))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "PDF to Word" in full_text or len(doc.paragraphs) > 0


def test_docx_to_pdf(engine, temp_dir):
    docx_path = temp_dir / "sample.docx"
    create_sample_docx(docx_path, "ReportLab Fallback Test", "Paragraph inside DOCX.")
    out_pdf = temp_dir / "sample.pdf"

    config = ConversionConfig(target_format="PDF")
    result = engine.convert_single(docx_path, out_pdf, config)

    assert result.success is True
    assert out_pdf.is_file()
    assert out_pdf.stat().st_size > 0

    # Verify generated PDF is valid via PyMuPDF
    doc = fitz.open(str(out_pdf))
    assert len(doc) >= 1
    doc.close()


def test_docx_to_txt(engine, temp_dir):
    docx_path = temp_dir / "sample.docx"
    create_sample_docx(docx_path, "Text Export Test", "Body line for export.")
    out_txt = temp_dir / "exported.txt"

    config = ConversionConfig(target_format="TXT")
    result = engine.convert_single(docx_path, out_txt, config)

    assert result.success is True
    assert out_txt.is_file()
    content = out_txt.read_text(encoding="utf-8")
    assert "Text Export Test" in content
    assert "Body line for export." in content


def test_docx_to_html(engine, temp_dir):
    docx_path = temp_dir / "sample.docx"
    create_sample_docx(docx_path, "HTML Heading", "HTML Paragraph")
    out_html = temp_dir / "page.html"

    config = ConversionConfig(target_format="HTML")
    result = engine.convert_single(docx_path, out_html, config)

    assert result.success is True
    assert out_html.is_file()
    content = out_html.read_text(encoding="utf-8")
    assert "<h1>HTML Heading</h1>" in content
    assert "<p>HTML Paragraph</p>" in content
    assert "<table" in content


def test_docx_to_png_images(engine, temp_dir):
    docx_path = temp_dir / "sample.docx"
    create_sample_docx(docx_path, "Docx To Images", "Rendered through PDF.")
    out_png = temp_dir / "page.png"

    config = ConversionConfig(target_format="PNG")
    result = engine.convert_single(docx_path, out_png, config)

    assert result.success is True
    assert Path(result.output_path).is_file()


def test_corrupt_page_in_table_extraction_surfaces_error(engine, temp_dir, monkeypatch):
    """
    Item 2 verification: When an exception occurs during table extraction (previously swallowed
    by bare 'except Exception: pass'), it must NOT be silently skipped. It must surface
    as a failure in the ConversionResult with the error message preserved.
    """
    pdf_path = temp_dir / "table_doc.pdf"
    create_sample_pdf(pdf_path, "PDF with tables")
    out_csv = temp_dir / "extracted.csv"

    def fake_find_tables(self):
        raise RuntimeError("Simulated corrupt table structure on page")

    monkeypatch.setattr(fitz.Page, "find_tables", fake_find_tables)

    config = ConversionConfig(target_format="CSV")
    result = engine.convert_single(pdf_path, out_csv, config)

    assert result.success is False
    assert result.error_message is not None
    assert "Simulated corrupt table structure on page" in result.error_message or "Table extraction failed" in result.error_message


def test_docx_to_pdf_exhausted_fallback_surfaces_failure(engine, temp_dir, monkeypatch):
    """
    Item 2 verification: When docx->pdf fallbacks are exhausted (MS Word COM fails),
    the conversion must return success=False with a specific error message, not a silent pass.
    """
    docx_path = temp_dir / "test.docx"
    create_sample_docx(docx_path, "Title", "Body")
    out_pdf = temp_dir / "test_out.pdf"

    # Mock win32com dispatch to raise in convert_docx_to_pdf
    import sys
    monkeypatch.setattr(sys, "platform", "win32")
    import types
    fake_win32 = types.ModuleType("win32com.client")
    def fake_dispatch(app):
        raise RuntimeError("Simulated Word COM crash")
    fake_win32.DispatchEx = fake_dispatch
    monkeypatch.setitem(sys.modules, "win32com.client", fake_win32)
    if "win32com" in sys.modules:
        monkeypatch.setattr(sys.modules["win32com"], "client", fake_win32, raising=False)
    try:
        import win32com.client
        monkeypatch.setattr(win32com.client, "DispatchEx", fake_dispatch, raising=False)
    except Exception:
        pass

    result = engine.convert_single(docx_path, out_pdf, ConversionConfig(target_format="PDF"))
    assert result.success is False
    assert "Microsoft Word" in result.error_message
    assert not out_pdf.exists(), "No partial or corrupt output file must be left behind on failed DOCX->PDF"


def test_xlsx_to_pdf_exhausted_fallback_surfaces_failure(engine, temp_dir, monkeypatch):
    """
    Item 2 verification: When xlsx->pdf fallbacks are exhausted (MS Excel COM fails),
    the conversion must return success=False with a specific error message.
    """
    import pandas as pd
    xlsx_path = temp_dir / "test.xlsx"
    pd.DataFrame({"A": [1, 2], "B": [3, 4]}).to_excel(str(xlsx_path), index=False)
    out_pdf = temp_dir / "test_out.pdf"

    import sys
    monkeypatch.setattr("os.name", "nt")
    import types
    fake_win32 = types.ModuleType("win32com.client")
    def fake_dispatch(app):
        raise RuntimeError("Simulated Excel COM crash")
    fake_win32.DispatchEx = fake_dispatch
    monkeypatch.setitem(sys.modules, "win32com.client", fake_win32)
    if "win32com" in sys.modules:
        monkeypatch.setattr(sys.modules["win32com"], "client", fake_win32, raising=False)
    try:
        import win32com.client
        monkeypatch.setattr(win32com.client, "DispatchEx", fake_dispatch, raising=False)
    except Exception:
        pass

    result = engine.convert_single(xlsx_path, out_pdf, ConversionConfig(target_format="PDF"))
    assert result.success is False
    assert "Microsoft Excel" in result.error_message
    assert not out_pdf.exists(), "No partial or corrupt output file must be left behind on failed XLSX->PDF"


def test_pptx_to_pdf_exhausted_fallback_surfaces_failure(engine, temp_dir, monkeypatch):
    """
    Item 2 verification: When pptx->pdf fallbacks are exhausted (MS PowerPoint COM fails),
    the conversion must return success=False with a specific error message.
    """
    pptx_path = temp_dir / "test.pptx"
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(pptx_path))
    out_pdf = temp_dir / "test_out.pdf"

    import sys
    monkeypatch.setattr("os.name", "nt")
    import types
    fake_win32 = types.ModuleType("win32com.client")
    def fake_dispatch(app):
        raise RuntimeError("Simulated PowerPoint COM crash")
    fake_win32.DispatchEx = fake_dispatch
    monkeypatch.setitem(sys.modules, "win32com.client", fake_win32)
    if "win32com" in sys.modules:
        monkeypatch.setattr(sys.modules["win32com"], "client", fake_win32, raising=False)
    try:
        import win32com.client
        monkeypatch.setattr(win32com.client, "DispatchEx", fake_dispatch, raising=False)
    except Exception:
        pass

    result = engine.convert_single(pptx_path, out_pdf, ConversionConfig(target_format="PDF"))
    assert result.success is False
    assert "Microsoft PowerPoint" in result.error_message
    assert not out_pdf.exists(), "No partial or corrupt output file must be left behind on failed PPTX->PDF"


def test_end_to_end_table_extraction_failure_surfaces_page_number(engine, temp_dir, monkeypatch):
    """
    Item 2 verification: End-to-end table extraction failure must surface the exact
    page number in the ConversionResult.error_message returned by engine.convert_single().
    """
    pdf_path = temp_dir / "multi_page_table.pdf"
    doc = fitz.open()
    p1 = doc.new_page(width=595, height=842)
    p1.insert_text((72, 72), "Page 1 Content", fontsize=12)
    p2 = doc.new_page(width=595, height=842)
    p2.insert_text((72, 72), "Page 2 Corrupt Table", fontsize=12)
    doc.save(str(pdf_path))
    doc.close()

    out_csv = temp_dir / "out_tables.csv"

    orig_find_tables = fitz.Page.find_tables
    def patched_find_tables(page_self):
        if page_self.number == 1:  # 0-indexed: second page
            raise RuntimeError("Corrupt table structure")
        return orig_find_tables(page_self)

    monkeypatch.setattr(fitz.Page, "find_tables", patched_find_tables)

    result = engine.convert_single(pdf_path, out_csv, ConversionConfig(target_format="CSV"))

    assert result.success is False
    assert result.error_message is not None
    assert "page 2" in result.error_message.lower(), f"Expected 'page 2' in error message: {result.error_message}"


def _draw_table_on_page(page, headers, rows, top_left=(50, 50), col_w=80, row_h=30):
    x0, y0 = top_left
    ncols = len(headers)
    nrows = len(rows) + 1
    w, h = ncols * col_w, nrows * row_h
    s = page.new_shape()
    s.draw_rect(fitz.Rect(x0, y0, x0 + w, y0 + h))
    for r in range(1, nrows):
        s.draw_line(fitz.Point(x0, y0 + r * row_h), fitz.Point(x0 + w, y0 + r * row_h))
    for c in range(1, ncols):
        s.draw_line(fitz.Point(x0 + c * col_w, y0), fitz.Point(x0 + c * col_w, y0 + h))
    s.finish()
    s.commit()
    for c, head in enumerate(headers):
        page.insert_text((x0 + c * col_w + 5, y0 + row_h - 10), str(head), fontsize=10)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            page.insert_text((x0 + c * col_w + 5, y0 + (r + 2) * row_h - 10), str(val), fontsize=9)


def test_table_extraction_ground_truth_cases(engine, temp_dir):
    """
    Item 9 Ground Truth Tests:
    (a) Single table across 2 pages with identical columns -> concatenated seamlessly
    (b) Single table across 2 pages with repeated header -> header row in 2nd page dropped, concatenated
    (c) Two different tables with different column counts (3 cols and 5 cols) ->
        separate sheets for XLSX (and separate files for CSV), NO all-NaN columns!
        Includes printed before/after dump.
    (d) PDF with no tables -> ValueError
    (e) Gate check: engine.convert_single returns success=True, output_path exists and > 0 bytes.
    """
    import pandas as pd
    from image_converter.core.document_engine import DocumentConverter

    # -------------------------------------------------------------------------
    # Case (a): Single table across 2 pages with identical columns
    # -------------------------------------------------------------------------
    pdf_a = temp_dir / "table_a.pdf"
    doc_a = fitz.open()
    p1 = doc_a.new_page(width=500, height=400)
    _draw_table_on_page(p1, ["ID", "Name", "Score"], [["1", "Alice", "95"], ["2", "Bob", "88"]])
    p2 = doc_a.new_page(width=500, height=400)
    _draw_table_on_page(p2, ["ID", "Name", "Score"], [["3", "Charlie", "92"], ["4", "Diana", "97"]])
    doc_a.save(str(pdf_a))
    doc_a.close()

    out_csv_a = temp_dir / "table_a.csv"
    res_a = engine.convert_single(pdf_a, out_csv_a, ConversionConfig(target_format="CSV"))
    assert res_a.success is True
    assert out_csv_a.exists()
    df_a = pd.read_csv(out_csv_a)
    assert list(df_a.columns) == ["ID", "Name", "Score"]
    assert len(df_a) == 4
    assert list(df_a["Name"]) == ["Alice", "Bob", "Charlie", "Diana"]

    # -------------------------------------------------------------------------
    # Case (b): Single table across 2 pages with repeated header
    # -------------------------------------------------------------------------
    pdf_b = temp_dir / "table_b.pdf"
    doc_b = fitz.open()
    p1 = doc_b.new_page(width=500, height=400)
    _draw_table_on_page(p1, ["Item", "Qty", "Price"], [["Widget", "10", "5.00"]])
    p2 = doc_b.new_page(width=500, height=400)
    # Page 2 table includes duplicate header row as row 0
    _draw_table_on_page(p2, ["Item", "Qty", "Price"], [["Item", "Qty", "Price"], ["Gadget", "20", "12.50"]])
    doc_b.save(str(pdf_b))
    doc_b.close()

    out_csv_b = temp_dir / "table_b.csv"
    res_b = engine.convert_single(pdf_b, out_csv_b, ConversionConfig(target_format="CSV"))
    assert res_b.success is True
    assert out_csv_b.exists()
    df_b = pd.read_csv(out_csv_b)
    assert list(df_b.columns) == ["Item", "Qty", "Price"]
    # Duplicate header dropped, only 2 real data rows remain
    assert len(df_b) == 2
    assert list(df_b["Item"]) == ["Widget", "Gadget"]

    # -------------------------------------------------------------------------
    # Case (c): Two different tables with different column counts (3 cols vs 5 cols)
    # -------------------------------------------------------------------------
    pdf_c = temp_dir / "table_c_different.pdf"
    doc_c = fitz.open()
    p1 = doc_c.new_page(width=500, height=400)
    _draw_table_on_page(p1, ["A", "B", "C"], [["1", "2", "3"], ["4", "5", "6"]], col_w=80)
    p2 = doc_c.new_page(width=500, height=400)
    _draw_table_on_page(p2, ["V", "W", "X", "Y", "Z"], [["10", "20", "30", "40", "50"]], col_w=60)
    doc_c.save(str(pdf_c))
    doc_c.close()

    # Before/After Demonstration Dump for Case (c)
    # BEFORE: Blind pd.concat would merge columns producing all-NaN fields
    raw_doc = fitz.open(str(pdf_c))
    t1 = raw_doc[0].find_tables()[0].to_pandas()
    t2 = raw_doc[1].find_tables()[0].to_pandas()
    raw_doc.close()
    before_naive_concat = pd.concat([t1, t2], ignore_index=True)

    print("\n--- CASE (c) BEFORE / AFTER DUMP ---")
    print(">>> BEFORE (Naive pd.concat - Column count mismatch ignored):")
    print(f"Shape: {before_naive_concat.shape}")
    print(before_naive_concat)
    print(f"Total NaN cells: {before_naive_concat.isna().sum().sum()}")

    # AFTER: XLSX Export creates separate sheets with 0 NaN columns
    out_xlsx_c = temp_dir / "table_c.xlsx"
    res_c = engine.convert_single(pdf_c, out_xlsx_c, ConversionConfig(target_format="XLSX"))
    assert res_c.success is True
    assert out_xlsx_c.exists()
    assert out_xlsx_c.stat().st_size > 0

    with pd.ExcelFile(out_xlsx_c) as xl:
        sheet_names = xl.sheet_names
        assert len(sheet_names) == 2
        assert "Table_1" in sheet_names
        assert "Table_2" in sheet_names
        df_sheet1 = xl.parse("Table_1")
        df_sheet2 = xl.parse("Table_2")

    print("\n>>> AFTER (Shape-Validated Table Extraction):")
    print(f"Sheet Table_1 Shape: {df_sheet1.shape} (Cols: {list(df_sheet1.columns)})")
    print(df_sheet1)
    print(f"Sheet Table_2 Shape: {df_sheet2.shape} (Cols: {list(df_sheet2.columns)})")
    print(df_sheet2)
    print("--- END CASE (c) DUMP ---\n")

    # Invariant: Neither sheet contains all-NaN columns or corrupted schema
    assert df_sheet1.shape == (2, 3)
    assert list(df_sheet1.columns) == ["A", "B", "C"]
    assert df_sheet1.isna().sum().sum() == 0

    assert df_sheet2.shape == (1, 5)
    assert list(df_sheet2.columns) == ["V", "W", "X", "Y", "Z"]
    assert df_sheet2.isna().sum().sum() == 0

    # Also test CSV export for Case (c): primary written to table_c.csv, secondary to table_c_table2.csv
    out_csv_c = temp_dir / "table_c.csv"
    res_csv_c = engine.convert_single(pdf_c, out_csv_c, ConversionConfig(target_format="CSV"))
    assert res_csv_c.success is True
    assert out_csv_c.exists()
    df_csv1 = pd.read_csv(out_csv_c)
    assert df_csv1.shape == (2, 3)
    assert list(df_csv1.columns) == ["A", "B", "C"]

    out_csv2 = temp_dir / "table_c_table2.csv"
    assert out_csv2.exists()
    df_csv2 = pd.read_csv(out_csv2)
    assert df_csv2.shape == (1, 5)
    assert list(df_csv2.columns) == ["V", "W", "X", "Y", "Z"]

    # -------------------------------------------------------------------------
    # Case (d): PDF with no tables -> ValueError
    # -------------------------------------------------------------------------
    pdf_d = temp_dir / "no_tables.pdf"
    doc_d = fitz.open()
    p_d = doc_d.new_page(width=400, height=300)
    p_d.insert_text((50, 50), "This is a document with plain text and zero tables.")
    doc_d.save(str(pdf_d))
    doc_d.close()

    out_d = temp_dir / "out_no_tables.csv"
    res_d = engine.convert_single(pdf_d, out_d, ConversionConfig(target_format="CSV"))
    assert res_d.success is False
    assert "No structured tables found in PDF" in res_d.error_message
    assert not out_d.exists()


def test_table_extraction_headerless_continuation(engine, temp_dir):
    """
    Item 9b.2: Headerless continuation (common real-world case).
    Page 1: Header + 3 data rows.
    Page 2: 3 data rows with NO header (same column count).
    Ground truth assertion: All 6 data rows end up in one table under original header;
    Page 2's first data row is NOT consumed as a header.
    """
    import pandas as pd

    pdf_p = temp_dir / "headerless_continuation.pdf"
    doc = fitz.open()

    # Page 1: Header ["ID", "Val1", "Val2"] + 3 data rows
    p1 = doc.new_page(width=500, height=400)
    s1 = p1.new_shape()
    s1.draw_rect(fitz.Rect(50, 50, 290, 170))
    for r in range(1, 4): s1.draw_line(fitz.Point(50, 50 + r * 30), fitz.Point(290, 50 + r * 30))
    for c in range(1, 3): s1.draw_line(fitz.Point(50 + c * 80, 50), fitz.Point(50 + c * 80, 170))
    s1.finish(); s1.commit()
    p1.insert_text((55, 70), "ID"); p1.insert_text((135, 70), "Val1"); p1.insert_text((215, 70), "Val2")
    p1.insert_text((55, 100), "1"); p1.insert_text((135, 100), "10"); p1.insert_text((215, 100), "100")
    p1.insert_text((55, 130), "2"); p1.insert_text((135, 130), "20"); p1.insert_text((215, 130), "200")
    p1.insert_text((55, 160), "3"); p1.insert_text((135, 160), "30"); p1.insert_text((215, 160), "300")

    # Page 2: NO header, 3 data rows
    p2 = doc.new_page(width=500, height=400)
    s2 = p2.new_shape()
    s2.draw_rect(fitz.Rect(50, 50, 290, 140))
    for r in range(1, 3): s2.draw_line(fitz.Point(50, 50 + r * 30), fitz.Point(290, 50 + r * 30))
    for c in range(1, 3): s2.draw_line(fitz.Point(50 + c * 80, 50), fitz.Point(50 + c * 80, 140))
    s2.finish(); s2.commit()
    p2.insert_text((55, 70), "4"); p2.insert_text((135, 70), "40"); p2.insert_text((215, 70), "400")
    p2.insert_text((55, 100), "5"); p2.insert_text((135, 100), "50"); p2.insert_text((215, 100), "500")
    p2.insert_text((55, 130), "6"); p2.insert_text((135, 130), "60"); p2.insert_text((215, 130), "600")

    doc.save(str(pdf_p))
    doc.close()

    out_csv = temp_dir / "headerless_out.csv"
    res = engine.convert_single(pdf_p, out_csv, ConversionConfig(target_format="CSV"))
    assert res.success is True
    assert out_csv.exists()

    df = pd.read_csv(out_csv)
    # 1. Columns must match original header
    assert [c.strip() for c in df.columns] == ["ID", "Val1", "Val2"]
    # 2. Exactly 6 data rows must be preserved
    assert len(df) == 6, f"Expected 6 rows, got {len(df)}: {df}"
    # 3. Values must be 1 through 6 in order
    assert [str(v).strip() for v in df["ID"].tolist()] == ["1", "2", "3", "4", "5", "6"]


def test_table_extraction_non_consecutive_grouping_policy(engine, temp_dir):
    """
    Item 9b.3: Non-consecutive tables policy.
    Page 1: 3-col table.
    Page 2: 5-col table.
    Page 3: 3-col table.
    Policy: Intervening incompatible table prevents grouping across pages.
    Result: 3 separate tables/sheets in document order (Table_1, Table_2, Table_3).
    """
    import pandas as pd
    from image_converter.core.document_engine import DocumentConverter

    def _draw_tbl(p, ncols, nrows, top_left=(50, 50), col_w=60, row_h=30):
        w, h = ncols * col_w, nrows * row_h
        s = p.new_shape()
        s.draw_rect(fitz.Rect(50, 50, 50 + w, 50 + h))
        for r in range(1, nrows): s.draw_line(fitz.Point(50, 50 + r * row_h), fitz.Point(50 + w, 50 + r * row_h))
        for c in range(1, ncols): s.draw_line(fitz.Point(50 + c * col_w, 50), fitz.Point(50 + c * col_w, 50 + h))
        s.finish(); s.commit()
        for c in range(ncols):
            p.insert_text((55 + c * col_w, 70), f"H{c+1}", fontsize=10)
            p.insert_text((55 + c * col_w, 100), f"V{c+1}", fontsize=9)

    pdf_p = temp_dir / "non_consec.pdf"
    doc = fitz.open()
    p1 = doc.new_page(width=500, height=400); _draw_tbl(p1, 3, 2)
    p2 = doc.new_page(width=500, height=400); _draw_tbl(p2, 5, 2)
    p3 = doc.new_page(width=500, height=400); _draw_tbl(p3, 3, 2)
    doc.save(str(pdf_p)); doc.close()

    out_xlsx = temp_dir / "non_consec.xlsx"
    de = DocumentConverter()
    de.convert_pdf_to_tables(pdf_p, out_xlsx, target_format="XLSX")

    with pd.ExcelFile(out_xlsx) as xl:
        sheet_names = xl.sheet_names
        # Policy assertion: 3 separate sheets, NOT merged into 2
        assert sheet_names == ["Table_1", "Table_2", "Table_3"]
        assert xl.parse("Table_1").shape[1] == 3
        assert xl.parse("Table_2").shape[1] == 5
        assert xl.parse("Table_3").shape[1] == 3


def test_table_extraction_single_page_failure_skipped(engine, temp_dir, caplog, monkeypatch):
    """
    Item 9b.4: If a table on a single page inside a multi-page document fails/corrupts,
    it must be skipped with a logged warning naming the page, and remaining pages still extract.
    """
    import logging
    pdf_p = temp_dir / "page_error.pdf"
    doc = fitz.open()
    # Page 1: valid table
    p1 = doc.new_page(width=500, height=400)
    _draw_table_on_page(p1, ["ColA", "ColB"], [["1", "2"]])

    # Page 2: will trigger error
    p2 = doc.new_page(width=500, height=400)
    p2.insert_text((50, 50), "Corrupt Page")

    # Page 3: valid table
    p3 = doc.new_page(width=500, height=400)
    _draw_table_on_page(p3, ["ColA", "ColB"], [["3", "4"]])

    doc.save(str(pdf_p)); doc.close()

    orig_find_tables = fitz.Page.find_tables
    def patched_find_tables(page_self):
        if page_self.number == 1:  # Page 2 (0-indexed 1)
            raise RuntimeError("Corrupt page structure")
        return orig_find_tables(page_self)

    monkeypatch.setattr(fitz.Page, "find_tables", patched_find_tables)

    out_csv = temp_dir / "page_error_out.csv"
    with caplog.at_level(logging.WARNING):
        res = engine.convert_single(pdf_p, out_csv, ConversionConfig(target_format="CSV"))

    assert res.success is True
    assert out_csv.exists()
    assert "extracted 2 of 3 pages; page(s) 2 failed" in res.error_message
    assert "page_error_out_table2.csv" in res.error_message

    # Logged warning verified naming page 2
    warning_found = any("page 2" in record.message.lower() and "skipping" in record.message.lower() for record in caplog.records)
    assert warning_found, f"Expected logged warning naming page 2, got: {[r.message for r in caplog.records]}"


def test_table_extraction_digit_bearing_headers_stay_separate(engine, temp_dir):
    """
    Item 9 (C1.b): New table with digit-bearing headers (e.g. years '2021', '2022', '2023')
    even with the same column count must stay separate when not an immediate continuation,
    not be treated as a continuation row due to containing digits.
    """
    import pandas as pd

    pdf_p = temp_dir / "digit_headers.pdf"
    doc = fitz.open()

    # Page 1: 3-col table with string headers
    p1 = doc.new_page(width=500, height=400)
    _draw_table_on_page(p1, ["Metric", "Target", "Status"], [["Sales", "100", "OK"]])

    # Page 2: Intervening text page (not an immediate page continuation)
    p2 = doc.new_page(width=500, height=400)
    p2.insert_text((50, 50), "Intervening chapter text without tables")

    # Page 3: 3-col table with digit-bearing headers ['2021', '2022', '2023']
    p3 = doc.new_page(width=500, height=400)
    _draw_table_on_page(p3, ["2021", "2022", "2023"], [["10", "20", "30"]])

    doc.save(str(pdf_p)); doc.close()

    out_xlsx = temp_dir / "digit_headers.xlsx"
    de = DocumentConverter()
    de.convert_pdf_to_tables(pdf_p, out_xlsx, target_format="XLSX")

    with pd.ExcelFile(out_xlsx) as xl:
        assert xl.sheet_names == ["Table_1", "Table_2"]
        df1 = xl.parse("Table_1")
        df2 = xl.parse("Table_2")
        assert list(df1.columns) == ["Metric", "Target", "Status"]
        assert [str(c) for c in df2.columns] == ["2021", "2022", "2023"]


def test_table_extraction_different_x_geometry_stays_separate(engine, temp_dir):
    """
    Item 9 (C1.c): Tables on consecutive pages with the SAME column count but DIFFERENT
    x-geometry (> 5 pt difference in column boundaries) must stay separate.
    """
    import pandas as pd

    pdf_p = temp_dir / "diff_geom.pdf"
    doc = fitz.open()

    # Page 1: 3-column table with cols at [50, 130, 210, 290] (widths: 80, 80, 80)
    p1 = doc.new_page(width=500, height=400)
    s1 = p1.new_shape()
    s1.draw_rect(fitz.Rect(50, 50, 290, 110))
    s1.draw_line(fitz.Point(50, 80), fitz.Point(290, 80))
    s1.draw_line(fitz.Point(130, 50), fitz.Point(130, 110))
    s1.draw_line(fitz.Point(210, 50), fitz.Point(210, 110))
    s1.finish(); s1.commit()
    p1.insert_text((55, 70), "ColA"); p1.insert_text((135, 70), "ColB"); p1.insert_text((215, 70), "ColC")
    p1.insert_text((55, 100), "R1A"); p1.insert_text((135, 100), "R1B"); p1.insert_text((215, 100), "R1C")

    # Page 2: 3-column table with cols at [50, 180, 310, 440] (widths: 130, 130, 130) -> >5pt diff!
    p2 = doc.new_page(width=500, height=400)
    s2 = p2.new_shape()
    s2.draw_rect(fitz.Rect(50, 50, 440, 110))
    s2.draw_line(fitz.Point(50, 80), fitz.Point(440, 80))
    s2.draw_line(fitz.Point(180, 50), fitz.Point(180, 110))
    s2.draw_line(fitz.Point(310, 50), fitz.Point(310, 110))
    s2.finish(); s2.commit()
    p2.insert_text((55, 70), "ColA"); p2.insert_text((185, 70), "ColB"); p2.insert_text((315, 70), "ColC")
    p2.insert_text((55, 100), "R2A"); p2.insert_text((185, 100), "R2B"); p2.insert_text((315, 100), "R2C")

    doc.save(str(pdf_p)); doc.close()

    out_xlsx = temp_dir / "diff_geom.xlsx"
    de = DocumentConverter()
    de.convert_pdf_to_tables(pdf_p, out_xlsx, target_format="XLSX")

    with pd.ExcelFile(out_xlsx) as xl:
        # Must NOT merge because column x-boundaries differ by > 5 pt
        assert xl.sheet_names == ["Table_1", "Table_2"]
        df1 = xl.parse("Table_1")
        df2 = xl.parse("Table_2")
        assert df1.iloc[0].tolist() == ["R1A", "R1B", "R1C"]
        assert df2.iloc[0].tolist() == ["R2A", "R2B", "R2C"]


def test_table_extraction_header_normalization(engine, temp_dir):
    """
    Item 9b.5: Header normalization handles whitespace and case differences ('Amount' vs 'amount ').
    Rule: Matching is case-insensitive and stripped. The canonical header from page 1 is preserved.
    """
    import pandas as pd

    pdf_p = temp_dir / "norm_header.pdf"
    doc = fitz.open()

    p1 = doc.new_page(width=500, height=400)
    _draw_table_on_page(p1, ["Amount", "Status"], [["$100", "Paid"]])

    p2 = doc.new_page(width=500, height=400)
    # Case and whitespace difference: 'amount ' vs 'STATUS'
    _draw_table_on_page(p2, ["amount ", "STATUS"], [["$200", "Pending"]])

    doc.save(str(pdf_p)); doc.close()

    out_csv = temp_dir / "norm_header_out.csv"
    res = engine.convert_single(pdf_p, out_csv, ConversionConfig(target_format="CSV"))
    assert res.success is True

    df = pd.read_csv(out_csv)
    # Assert merged into 1 table with canonical casing 'Amount', 'Status'
    assert list(df.columns) == ["Amount", "Status"]
    assert len(df) == 2
    assert df["Amount"].tolist() == ["$100", "$200"]


def test_table_extraction_csv_side_files_cleanup_and_atomic_failure(engine, temp_dir, monkeypatch):
    """
    Item 9b.6 & C3: CSV side files policy:
    1. Clear stale {stem}_table*.csv files recorded in manifest from earlier runs.
    2. User file {stem}_table_notes.csv MUST survive.
    3. Atomic write with mid-write failure snapshot diff (0 partial files left).
    """
    import json
    import pandas as pd

    # Create user file that must survive
    user_notes = temp_dir / "side_test_table_notes.csv"
    user_notes.write_text("user notes content must survive")

    # Create stale side files and manifest from a hypothetical prior run
    stale1 = temp_dir / "side_test_table1.csv"; stale1.write_text("stale1")
    stale2 = temp_dir / "side_test_table2.csv"; stale2.write_text("stale2")
    stale9 = temp_dir / "side_test_table9.csv"; stale9.write_text("stale9")
    manifest = temp_dir / ".side_test_tables_manifest.json"
    manifest.write_text(json.dumps(["side_test.csv", "side_test_table1.csv", "side_test_table2.csv", "side_test_table9.csv"]))

    # Multi-table PDF (2 incompatible tables)
    pdf_p = temp_dir / "side_test.pdf"
    doc = fitz.open()
    p1 = doc.new_page(width=500, height=400)
    _draw_table_on_page(p1, ["A", "B"], [["1", "2"]])

    p2 = doc.new_page(width=500, height=400)
    _draw_table_on_page(p2, ["X", "Y", "Z"], [["10", "20", "30"]])
    doc.save(str(pdf_p)); doc.close()

    out_csv = temp_dir / "side_test.csv"
    res = engine.convert_single(pdf_p, out_csv, ConversionConfig(target_format="CSV"))
    assert res.success is True

    # Stale file stale9 must be gone
    assert not stale9.exists(), "Stale side file recorded in manifest was not cleared"
    assert (temp_dir / "side_test_table1.csv").exists()
    assert (temp_dir / "side_test_table2.csv").exists()

    # User file MUST survive
    assert user_notes.exists(), "User file side_test_table_notes.csv was improperly deleted!"
    assert user_notes.read_text() == "user notes content must survive"

    # Test atomic mid-write failure snapshot diff
    fail_csv = temp_dir / "atomic_fail.csv"
    files_before = set(temp_dir.iterdir())

    # Simulate write failure on side file
    orig_to_csv = pd.DataFrame.to_csv
    call_count = 0
    def failing_to_csv(self, *args, **kwargs):
        nonlocal call_count
        call_count += 1
        if call_count == 2:
            raise OSError("Simulated disk failure during CSV write")
        return orig_to_csv(self, *args, **kwargs)

    monkeypatch.setattr(pd.DataFrame, "to_csv", failing_to_csv)

    res_fail = engine.convert_single(pdf_p, fail_csv, ConversionConfig(target_format="CSV"))
    assert res_fail.success is False

    files_after = set(temp_dir.iterdir())
    # Snapshot diff: zero new or lingering files created
    new_files = files_after - files_before
    assert len(new_files) == 0, f"Mid-write failure left partial files: {new_files}"


