"""
Tests for metadata inspection and stripping across image, document, and spreadsheet formats.
"""

from pathlib import Path
import zipfile
import docx
import fitz
import openpyxl
from PIL import Image
import pytest

from image_converter.cli import main
from image_converter.core.engine import UniversalConverterEngine, ConversionConfig
from image_converter.core.metadata_engine import (
    get_detailed_metadata,
    strip_file_metadata,
    strip_image_metadata,
    strip_pdf_metadata,
    strip_docx_metadata,
    strip_xlsx_metadata,
    strip_pptx_metadata,
    strip_odt_metadata,
)

try:
    import pptx
except ImportError:
    pptx = None


def test_image_metadata_strip(tmp_path: Path):
    """Test creating a JPEG with EXIF tags, inspecting it, and stripping all metadata."""
    img_path = tmp_path / "photo_with_exif.jpg"
    clean_path = tmp_path / "photo_clean.jpg"

    # Create image with EXIF
    img = Image.new("RGB", (80, 80), color=(100, 150, 200))
    exif = img.getexif()
    exif[0x010E] = "Secret Photo Description"  # ImageDescription
    exif[0x010F] = "Camera Brand X"             # Make
    exif[0x0110] = "Model 123"                  # Model
    exif[0x0132] = "2026:09:18 10:00:00"        # DateTime
    img.save(img_path, exif=exif)

    # 1. Inspect
    meta_before = get_detailed_metadata(img_path)
    assert meta_before["has_metadata"] is True
    assert meta_before["fields_count"] >= 3
    assert "Make" in meta_before["fields"]
    assert any("device" in w.lower() or "author" in w.lower() for w in meta_before["warnings"])

    # 2. Strip
    res = strip_file_metadata(img_path, clean_path)
    assert res["success"] is True
    assert clean_path.exists()

    # 3. Verify clean
    meta_after = get_detailed_metadata(clean_path)
    assert meta_after["has_metadata"] is False
    assert meta_after["fields_count"] == 0
    assert len(meta_after["warnings"]) == 0

    # Image should still be completely valid and readable
    with Image.open(clean_path) as verified_img:
        assert verified_img.size == (80, 80)


def test_pdf_metadata_strip(tmp_path: Path):
    """Test PDF creation with author, title, and XMP XML metadata, followed by stripping."""
    pdf_path = tmp_path / "document_tracked.pdf"
    clean_path = tmp_path / "document_clean.pdf"

    # Create PDF with metadata
    doc = fitz.open()
    page = doc.new_page(width=300, height=300)
    page.insert_text((50, 50), "Classified Content", fontsize=14)
    doc.set_metadata({
        "title": "Confidential Report",
        "author": "Dr. John Doe",
        "subject": "Top Secret Data",
        "creator": "Universal Converter",
        "producer": "PyMuPDF Test",
    })
    doc.set_xml_metadata("<x:xmpmeta><rdf:RDF><rdf:Description><dc:creator>John Doe</dc:creator></rdf:Description></rdf:RDF></x:xmpmeta>")
    doc.save(str(pdf_path))
    doc.close()

    # 1. Inspect
    meta_before = get_detailed_metadata(pdf_path)
    assert meta_before["has_metadata"] is True
    assert "Author" in meta_before["fields"]
    assert meta_before["fields"]["Author"] == "Dr. John Doe"
    assert any("author" in w.lower() for w in meta_before["warnings"])

    # 2. Strip
    res = strip_pdf_metadata(pdf_path, clean_path)
    assert len(res) > 0
    assert clean_path.exists()

    # 3. Verify clean
    meta_after = get_detailed_metadata(clean_path)
    assert meta_after["has_metadata"] is False

    clean_doc = fitz.open(str(clean_path))
    assert clean_doc.metadata.get("author") in (None, "")
    assert clean_doc.metadata.get("title") in (None, "")
    assert clean_doc.get_xml_metadata() in (None, "")
    assert len(clean_doc) == 1
    clean_doc.close()


def test_docx_metadata_strip(tmp_path: Path):
    """Test DOCX core properties inspection and stripping."""
    docx_path = tmp_path / "memo.docx"
    clean_path = tmp_path / "memo_clean.docx"

    # Create DOCX with core properties
    doc = docx.Document()
    doc.add_heading("Company Memo", level=1)
    doc.add_paragraph("This memo contains sensitive internal discussion.")
    doc.core_properties.author = "Executive Alice"
    doc.core_properties.title = "Internal Strategy"
    doc.core_properties.comments = "Reviewed by Board"
    doc.core_properties.last_modified_by = "Alice Editor"
    doc.save(str(docx_path))

    # 1. Inspect
    meta_before = get_detailed_metadata(docx_path)
    assert meta_before["has_metadata"] is True
    assert meta_before["fields"].get("Author") == "Executive Alice"
    assert meta_before["fields"].get("Title") == "Internal Strategy"

    # 2. Strip
    res = strip_docx_metadata(docx_path, clean_path)
    assert len(res) > 0
    assert clean_path.exists()

    # 3. Verify clean
    reopened = docx.Document(str(clean_path))
    assert reopened.core_properties.author in (None, "")
    assert reopened.core_properties.title in (None, "")
    assert reopened.core_properties.comments in (None, "")
    assert reopened.core_properties.last_modified_by in (None, "")
    assert len(reopened.paragraphs) >= 1


def test_xlsx_metadata_strip(tmp_path: Path):
    """Test XLSX creator/properties inspection and stripping."""
    xlsx_path = tmp_path / "budget.xlsx"
    clean_path = tmp_path / "budget_clean.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q3 Revenue"
    ws.append(["Item", "Amount"])
    ws.append(["Server Cost", 1250])
    wb.properties.creator = "Accountant Bob"
    wb.properties.title = "Q3 Financials"
    wb.properties.lastModifiedBy = "Auditor Dave"
    wb.save(str(xlsx_path))

    # 1. Inspect
    meta_before = get_detailed_metadata(xlsx_path)
    assert meta_before["has_metadata"] is True
    assert meta_before["fields"].get("Creator") == "Accountant Bob"

    # 2. Strip
    res = strip_xlsx_metadata(xlsx_path, clean_path)
    assert len(res) > 0
    assert clean_path.exists()

    # 3. Verify clean
    reopened_wb = openpyxl.load_workbook(str(clean_path))
    assert reopened_wb.properties.creator in (None, "")
    assert reopened_wb.properties.title in (None, "")
    assert reopened_wb.properties.lastModifiedBy in (None, "")
    assert reopened_wb.active["A1"].value == "Item"
    assert reopened_wb.active["B2"].value == 1250


def test_pptx_metadata_strip(tmp_path: Path):
    """Test PPTX core properties inspection and stripping."""
    if pptx is None:
        pytest.skip("python-pptx not available")

    pptx_path = tmp_path / "slides.pptx"
    clean_path = tmp_path / "slides_clean.pptx"

    prs = pptx.Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Pitch Deck"
    prs.core_properties.author = "Founder Charlie"
    prs.core_properties.title = "Investor Pitch"
    prs.save(str(pptx_path))

    # 1. Inspect
    meta_before = get_detailed_metadata(pptx_path)
    assert meta_before["has_metadata"] is True
    assert meta_before["fields"].get("Author") == "Founder Charlie"

    # 2. Strip
    res = strip_pptx_metadata(pptx_path, clean_path)
    assert len(res) > 0

    # 3. Verify clean
    reopened_prs = pptx.Presentation(str(clean_path))
    assert reopened_prs.core_properties.author in (None, "")
    assert reopened_prs.core_properties.title in (None, "")


def test_odt_metadata_strip(tmp_path: Path):
    """Test OpenDocument ODT meta.xml stripping."""
    odt_path = tmp_path / "sample.odt"
    clean_path = tmp_path / "sample_clean.odt"

    # Create a synthetic minimal ODT zip archive
    meta_xml_content = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<office:document-meta xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0">\n'
        '<office:meta>\n'
        '<meta:generator>LibreOffice/7.0</meta:generator>\n'
        '<dc:title>Test ODT</dc:title>\n'
        '<dc:creator>David Author</dc:creator>\n'
        '</office:meta>\n'
        '</office:document-meta>'
    )
    content_xml = '<?xml version="1.0" encoding="UTF-8"?><office:document-content></office:document-content>'

    with zipfile.ZipFile(str(odt_path), "w") as z:
        z.writestr("meta.xml", meta_xml_content)
        z.writestr("content.xml", content_xml)

    # 1. Inspect
    meta_before = get_detailed_metadata(odt_path)
    assert meta_before["has_metadata"] is True
    assert meta_before["fields"].get("Creator") == "David Author"

    # 2. Strip
    res = strip_odt_metadata(odt_path, clean_path)
    assert len(res) > 0

    # 3. Verify clean
    with zipfile.ZipFile(str(clean_path), "r") as z:
        cleaned_meta = z.read("meta.xml").decode("utf-8")
        assert "David Author" not in cleaned_meta
        assert "LibreOffice/7.0" not in cleaned_meta


def test_engine_strip_metadata_single(tmp_path: Path):
    """Test UniversalConverterEngine with strip_metadata configuration."""
    engine = UniversalConverterEngine()

    # 1. Create PDF with metadata
    pdf_path = tmp_path / "engine_input.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.set_metadata({"author": "Secret User", "title": "Confidential Engine Test"})
    doc.save(str(pdf_path))
    doc.close()

    # 2. Test engine.strip_metadata() directly
    clean_direct = tmp_path / "engine_direct_clean.pdf"
    res = engine.strip_metadata(pdf_path, clean_direct)
    assert res.success is True
    assert clean_direct.exists()
    d = fitz.open(str(clean_direct))
    assert d.metadata.get("author") in (None, "")
    d.close()

    # 3. Test engine.convert_single with target_format="STRIP_METADATA"
    clean_target = tmp_path / "engine_target_clean.pdf"
    cfg = ConversionConfig(target_format="STRIP_METADATA")
    res2 = engine.convert_single(pdf_path, clean_target, cfg)
    assert res2.success is True
    assert clean_target.exists()


def test_cli_inspect_metadata(tmp_path: Path, capsys):
    """Test CLI --inspect-metadata option."""
    pdf_path = tmp_path / "cli_inspect.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.set_metadata({"author": "Inspector Gadget", "title": "Gadget Dossier"})
    doc.save(str(pdf_path))
    doc.close()

    exit_code = main(["--inspect-metadata", str(pdf_path)])
    assert exit_code == 0
    captured = capsys.readouterr().out
    assert "Metadata Inspection" in captured
    assert "Inspector Gadget" in captured
    assert "Gadget Dossier" in captured


def test_cli_clean_metadata(tmp_path: Path):
    """Test CLI --clean-metadata option."""
    pdf_path = tmp_path / "cli_clean_target.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.set_metadata({"author": "To Be Removed", "title": "Will Be Cleaned"})
    doc.save(str(pdf_path))
    doc.close()

    exit_code = main(["--clean-metadata", str(pdf_path)])
    assert exit_code == 0

    expected_clean_file = tmp_path / "cli_clean_target_clean.pdf"
    assert expected_clean_file.exists()

    meta = get_detailed_metadata(expected_clean_file)
    assert meta["has_metadata"] is False


def test_cli_conversion_with_strip_metadata(tmp_path: Path):
    """Test format conversion with --strip-metadata flag."""
    img_path = tmp_path / "input.jpg"
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    img = Image.new("RGB", (60, 60), color="green")
    exif = img.getexif()
    exif[0x010F] = "Camera Brand Z"
    img.save(img_path, exif=exif)

    exit_code = main(["-i", str(img_path), "-f", "PNG", "-o", str(out_dir), "--strip-metadata"])
    assert exit_code == 0

    out_file = out_dir / "input.png"
    assert out_file.exists()

    meta = get_detailed_metadata(out_file)
    assert meta["has_metadata"] is False
