"""
Comprehensive QA Audit Test Suite for Universal File Converter.

Covers 14 sections:
  1. ConversionConfig   2. ConversionResult   3. Image conversions
  4. Document conversions   5. Data/spreadsheet   6. Security
  7. Edge cases   8. Metadata   9. OCR   10. Batch/concurrency
  11. CLI   12. GUI widgets   13. Logging   14. Integration
"""

from __future__ import annotations

import json
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from typing import List
from unittest.mock import patch

import docx
import fitz
import openpyxl
import pandas as pd
import pytest
from PIL import Image

from image_converter.core.engine import (
    ConversionConfig,
    ConversionResult,
    FORMAT_EXTENSIONS,
    ImageConverterEngine,
    SUPPORTED_INPUT_FORMATS,
    SUPPORTED_OUTPUT_FORMATS,
    UniversalConverterEngine,
    _HEIF_AVAILABLE,
    get_file_metadata,
    get_image_metadata,
    get_supported_input_extensions,
    get_supported_output_formats,
)
from image_converter.core.security import (
    SecurityError,
    sniff_file_type,
    validate_input_file,
    validate_output_path,
    validate_zip_container,
)
from image_converter.core.metadata_engine import (
    get_detailed_metadata,
    strip_file_metadata,
    strip_image_metadata,
    strip_pdf_metadata,
    strip_docx_metadata,
    strip_xlsx_metadata,
    strip_odt_metadata,
)
from image_converter.core.ocr_engine import (
    is_ocr_available,
    get_ocr_status_message,
    pdf_needs_ocr,
    ocr_image_to_text,
)
from image_converter.core.document_engine import (
    DocumentConverter,
    open_pdf_with_password,
    check_docx_encryption,
    get_pdf_metadata,
    get_docx_metadata,
)
from image_converter.core.data_engine import (
    DataConverter,
    get_csv_metadata,
    get_excel_metadata,
)
from image_converter.core.presentation_engine import (
    PresentationConverter,
    get_pptx_metadata,
)
from image_converter.core.rich_doc_engine import (
    RichDocumentConverter,
    get_odt_metadata,
    get_rtf_metadata,
)
from image_converter.core.logging_config import setup_logger, get_logger
from image_converter.cli import main, parse_color, collect_convertible_files

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from odf import opendocument, text as odf_text
except ImportError:
    opendocument = None


# ============================================================================
# FIXTURES
# ============================================================================

@pytest.fixture
def engine():
    return ImageConverterEngine()


# ============================================================================
# HELPERS
# ============================================================================

def _make_image(path: Path, w=80, h=80, color="red", mode="RGB") -> Path:
    img = Image.new(mode, (w, h), color=color)
    img.save(path)
    return path


def _make_image_with_exif(path: Path) -> Path:
    img = Image.new("RGB", (80, 80), color="blue")
    exif = img.getexif()
    exif[0x010F] = "TestMake"
    exif[0x0110] = "TestModel"
    exif[0x0132] = "2026:01:01 00:00:00"
    img.save(path, exif=exif)
    return path


def _make_pdf(path: Path, text="Hello QA Audit") -> Path:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 72), text, fontsize=14)
    doc.save(str(path))
    doc.close()
    return path


def _make_pdf_with_meta(path: Path) -> Path:
    doc = fitz.open()
    doc.new_page()
    doc.set_metadata({"author": "AuditAuthor", "title": "AuditTitle"})
    doc.save(str(path))
    doc.close()
    return path


def _make_docx(path: Path, title="QA Heading", body="QA body text.") -> Path:
    d = docx.Document()
    d.add_heading(title, level=1)
    d.add_paragraph(body)
    d.save(str(path))
    return path


def _make_docx_with_meta(path: Path) -> Path:
    d = docx.Document()
    d.add_paragraph("meta test")
    d.core_properties.author = "AuditDocxAuthor"
    d.core_properties.title = "AuditDocxTitle"
    d.save(str(path))
    return path


def _make_csv(path: Path) -> Path:
    pd.DataFrame({"Name": ["Alice", "Bob"], "Score": [95, 87]}).to_csv(str(path), index=False)
    return path


def _make_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Col1", "Col2"])
    ws.append(["A", 10])
    ws.append(["B", 20])
    wb.save(str(path))
    return path


def _make_xlsx_with_meta(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Data"])
    wb.properties.creator = "AuditCreator"
    wb.properties.title = "AuditXlsxTitle"
    wb.save(str(path))
    return path


def _make_pptx(path: Path) -> Path:
    if Presentation is None:
        pytest.skip("python-pptx not installed")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "QA Audit Slide"
    prs.save(str(path))
    return path


def _make_odt(path: Path) -> Path:
    if opendocument is None:
        pytest.skip("odfpy not installed")
    doc = opendocument.OpenDocumentText()
    doc.text.addElement(odf_text.P(text="QA ODT paragraph"))
    doc.save(str(path))
    return path


def _make_rtf(path: Path) -> Path:
    rtf = r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Calibri;}}\f0\fs24 QA RTF Content\par}"
    path.write_text(rtf, encoding="utf-8")
    return path


# ============================================================================
# SECTION 1: ConversionConfig Audit
# ============================================================================

class TestConversionConfig:

    def test_defaults(self):
        """All default values match expected."""
        c = ConversionConfig()
        assert c.target_format == "JPG"
        assert c.quality == 90
        assert c.lossless is False
        assert c.preserve_metadata is True
        assert c.strip_metadata is False
        assert c.auto_orient is True
        assert c.resize_mode == "none"
        assert c.resize_percent == 100.0
        assert c.custom_width is None
        assert c.custom_height is None
        assert c.keep_aspect_ratio is True
        assert c.background_color == (255, 255, 255)
        assert c.dpi == 150
        assert c.sheet_name is None
        assert c.password is None
        assert c.enable_ocr is False
        assert c.max_workers is None
        assert c.max_file_size == 500 * 1024 * 1024

    def test_strip_metadata_disables_preserve(self):
        """strip_metadata=True should set preserve_metadata=False."""
        c = ConversionConfig(strip_metadata=True)
        assert c.strip_metadata is True
        assert c.preserve_metadata is False

    def test_preserve_false_enables_strip(self):
        """preserve_metadata=False should set strip_metadata=True."""
        c = ConversionConfig(preserve_metadata=False)
        assert c.strip_metadata is True
        assert c.preserve_metadata is False

    def test_quality_accepts_any_int(self):
        """Quality is stored as-is (clamped at save time)."""
        assert ConversionConfig(quality=1).quality == 1
        assert ConversionConfig(quality=100).quality == 100
        assert ConversionConfig(quality=200).quality == 200

    def test_resize_mode_percentage(self):
        c = ConversionConfig(resize_mode="percentage", resize_percent=50.0)
        assert c.resize_mode == "percentage"
        assert c.resize_percent == 50.0

    def test_resize_mode_custom(self):
        c = ConversionConfig(resize_mode="custom", custom_width=200, custom_height=100)
        assert c.custom_width == 200
        assert c.custom_height == 100

    def test_all_format_keys_present(self):
        """FORMAT_EXTENSIONS has entries for all output formats."""
        for fmt in SUPPORTED_OUTPUT_FORMATS:
            assert fmt in FORMAT_EXTENSIONS, f"Missing FORMAT_EXTENSIONS key: {fmt}"

    def test_serialization_round_trip(self):
        """Config can be serialized to dict and reconstructed."""
        c = ConversionConfig(target_format="WEBP", quality=75, dpi=300, enable_ocr=True)
        d = {f.name: getattr(c, f.name) for f in c.__dataclass_fields__.values()}
        c2 = ConversionConfig(**d)
        assert c2.target_format == "WEBP"
        assert c2.quality == 75
        assert c2.dpi == 300
        assert c2.enable_ocr is True


# ============================================================================
# SECTION 2: ConversionResult Audit
# ============================================================================

class TestConversionResult:

    def test_success_result(self):
        r = ConversionResult(success=True, input_path="a.jpg", output_path="b.png",
                             input_format="JPEG", output_format="PNG")
        assert r.success is True
        assert r.error_message is None

    def test_failure_result(self):
        r = ConversionResult(success=False, input_path="a.jpg", error_message="bad file")
        assert r.success is False
        assert "bad file" in r.error_message

    def test_duration_positive(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.png")
        dst = tmp_path / "dst.jpg"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="JPG"))
        assert res.duration_seconds > 0

    def test_sizes_populated(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.png")
        dst = tmp_path / "dst.jpg"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="JPG"))
        assert res.input_size_bytes > 0
        assert res.output_size_bytes > 0

    def test_dimensions_tracked(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.png", w=120, h=60)
        dst = tmp_path / "dst.jpg"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="JPG"))
        assert res.input_dimensions == (120, 60)
        assert res.output_dimensions == (120, 60)


# ============================================================================
# SECTION 3: Image Conversion Matrix
# ============================================================================

class TestImageConversions:

    @pytest.mark.parametrize("target", ["JPG", "WEBP", "BMP", "TIFF", "GIF", "ICO", "PDF"])
    def test_png_to_formats(self, engine, tmp_path, target):
        """PNG -> various image formats."""
        src = _make_image(tmp_path / "src.png", mode="RGBA" if target != "PDF" else "RGB")
        dst = tmp_path / f"out{FORMAT_EXTENSIONS[target]}"
        res = engine.convert_single(src, dst, ConversionConfig(target_format=target))
        assert res.success is True, res.error_message
        assert Path(res.output_path).stat().st_size > 0

    @pytest.mark.parametrize("target", ["PNG", "WEBP", "BMP", "TIFF", "GIF", "ICO", "PDF"])
    def test_jpg_to_formats(self, engine, tmp_path, target):
        """JPG -> various image formats."""
        src = _make_image(tmp_path / "src.jpg")
        dst = tmp_path / f"out{FORMAT_EXTENSIONS[target]}"
        res = engine.convert_single(src, dst, ConversionConfig(target_format=target))
        assert res.success is True, res.error_message
        assert Path(res.output_path).stat().st_size > 0

    def test_webp_to_png(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.webp")
        dst = tmp_path / "out.png"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PNG"))
        assert res.success is True

    def test_bmp_to_jpg(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.bmp")
        dst = tmp_path / "out.jpg"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="JPG"))
        assert res.success is True

    def test_tiff_to_png(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.tiff")
        dst = tmp_path / "out.png"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PNG"))
        assert res.success is True

    def test_gif_to_jpg(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.gif")
        dst = tmp_path / "out.jpg"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="JPG"))
        assert res.success is True

    def test_alpha_compositing_to_jpg(self, engine, tmp_path):
        """Transparent RGBA -> JPG composites onto background color."""
        src = _make_image(tmp_path / "alpha.png", mode="RGBA", color=(255, 0, 0, 0))
        dst = tmp_path / "out.jpg"
        res = engine.convert_single(src, dst, ConversionConfig(
            target_format="JPG", background_color=(0, 0, 255)))
        assert res.success is True
        with Image.open(dst) as img:
            assert img.mode == "RGB"

    def test_resize_percentage(self, engine, tmp_path):
        """Resize by 50%."""
        src = _make_image(tmp_path / "src.png", w=200, h=100)
        dst = tmp_path / "out.png"
        res = engine.convert_single(src, dst, ConversionConfig(
            target_format="PNG", resize_mode="percentage", resize_percent=50.0))
        assert res.success is True
        assert res.output_dimensions == (100, 50)

    def test_resize_custom_keep_ratio(self, engine, tmp_path):
        """Custom resize with aspect ratio preservation."""
        src = _make_image(tmp_path / "src.png", w=200, h=100)
        dst = tmp_path / "out.png"
        res = engine.convert_single(src, dst, ConversionConfig(
            target_format="PNG", resize_mode="custom",
            custom_width=100, custom_height=100, keep_aspect_ratio=True))
        assert res.success is True
        assert res.output_dimensions == (100, 50)


# ============================================================================
# SECTION 4: Document Conversion Matrix
# ============================================================================

class TestDocumentConversions:

    def test_pdf_to_txt(self, engine, tmp_path):
        src = _make_pdf(tmp_path / "doc.pdf", "QA audit text extraction")
        dst = tmp_path / "doc.txt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="TXT"))
        assert res.success is True
        assert "QA audit text extraction" in dst.read_text(encoding="utf-8")

    def test_pdf_to_docx(self, engine, tmp_path):
        src = _make_pdf(tmp_path / "doc.pdf", "PDF to Word test")
        dst = tmp_path / "doc.docx"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="DOCX"))
        assert res.success is True
        assert dst.stat().st_size > 0

    def test_pdf_to_png(self, engine, tmp_path):
        src = _make_pdf(tmp_path / "doc.pdf")
        dst = tmp_path / "doc.png"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PNG"))
        assert res.success is True

    def test_docx_to_pdf(self, engine, tmp_path):
        src = _make_docx(tmp_path / "doc.docx")
        dst = tmp_path / "doc.pdf"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PDF"))
        assert res.success is True
        d = fitz.open(str(dst))
        assert len(d) >= 1
        d.close()

    def test_docx_to_txt(self, engine, tmp_path):
        src = _make_docx(tmp_path / "doc.docx", body="Extractable text.")
        dst = tmp_path / "doc.txt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="TXT"))
        assert res.success is True
        assert "Extractable text" in dst.read_text(encoding="utf-8")

    def test_docx_to_html(self, engine, tmp_path):
        src = _make_docx(tmp_path / "doc.docx", title="HTML Test")
        dst = tmp_path / "doc.html"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="HTML"))
        assert res.success is True
        assert "<h1>" in dst.read_text(encoding="utf-8").lower() or "<p>" in dst.read_text(encoding="utf-8").lower()

    @pytest.mark.office
    def test_docx_to_odt(self, engine, tmp_path):
        src = _make_docx(tmp_path / "doc.docx")
        dst = tmp_path / "doc.odt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="ODT"))
        assert res.success is True
        assert dst.stat().st_size > 0

    @pytest.mark.office
    def test_pptx_to_pdf(self, engine, tmp_path):
        src = _make_pptx(tmp_path / "pres.pptx")
        dst = tmp_path / "pres.pdf"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PDF"))
        assert res.success is True

    def test_pptx_to_txt(self, engine, tmp_path):
        src = _make_pptx(tmp_path / "pres.pptx")
        dst = tmp_path / "pres.txt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="TXT"))
        assert res.success is True
        assert "QA Audit Slide" in dst.read_text(encoding="utf-8")

    def test_odt_to_txt(self, engine, tmp_path):
        src = _make_odt(tmp_path / "doc.odt")
        dst = tmp_path / "doc.txt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="TXT"))
        assert res.success is True

    def test_rtf_to_txt(self, engine, tmp_path):
        src = _make_rtf(tmp_path / "doc.rtf")
        dst = tmp_path / "doc.txt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="TXT"))
        assert res.success is True
        assert "QA RTF Content" in dst.read_text(encoding="utf-8")

    def test_txt_to_pdf(self, engine, tmp_path):
        src = tmp_path / "input.txt"
        src.write_text("Plain text to PDF audit.", encoding="utf-8")
        dst = tmp_path / "out.pdf"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PDF"))
        assert res.success is True

    def test_txt_to_docx(self, engine, tmp_path):
        src = tmp_path / "input.txt"
        src.write_text("Text to DOCX.", encoding="utf-8")
        dst = tmp_path / "out.docx"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="DOCX"))
        assert res.success is True

    def test_txt_to_pptx(self, engine, tmp_path):
        src = tmp_path / "input.txt"
        src.write_text("=== Slide 1 ===\nContent", encoding="utf-8")
        dst = tmp_path / "out.pptx"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PPTX"))
        assert res.success is True


# ============================================================================
# SECTION 5: Data/Spreadsheet Conversion Matrix
# ============================================================================

class TestDataConversions:

    def test_csv_to_xlsx(self, engine, tmp_path):
        src = _make_csv(tmp_path / "data.csv")
        dst = tmp_path / "data.xlsx"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="XLSX"))
        assert res.success is True
        wb = openpyxl.load_workbook(str(dst))
        assert wb.active.cell(1, 1).value == "Name"
        wb.close()

    def test_csv_to_pdf(self, engine, tmp_path):
        src = _make_csv(tmp_path / "data.csv")
        dst = tmp_path / "data.pdf"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PDF"))
        assert res.success is True

    def test_csv_to_json(self, engine, tmp_path):
        src = _make_csv(tmp_path / "data.csv")
        dst = tmp_path / "data.json"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="JSON"))
        assert res.success is True
        data = json.loads(dst.read_text(encoding="utf-8"))
        assert isinstance(data, list)
        assert len(data) == 2
        assert data[0]["Name"] == "Alice"

    def test_csv_to_html(self, engine, tmp_path):
        src = _make_csv(tmp_path / "data.csv")
        dst = tmp_path / "data.html"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="HTML"))
        assert res.success is True
        assert "<table" in dst.read_text(encoding="utf-8")

    def test_csv_to_txt(self, engine, tmp_path):
        src = _make_csv(tmp_path / "data.csv")
        dst = tmp_path / "data.txt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="TXT"))
        assert res.success is True
        assert "Alice" in dst.read_text(encoding="utf-8")

    def test_xlsx_to_csv(self, engine, tmp_path):
        src = _make_xlsx(tmp_path / "data.xlsx")
        dst = tmp_path / "data.csv"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="CSV"))
        assert res.success is True
        df = pd.read_csv(str(dst))
        assert "Col1" in df.columns
        assert len(df) == 2

    def test_xlsx_to_pdf(self, engine, tmp_path):
        src = _make_xlsx(tmp_path / "data.xlsx")
        dst = tmp_path / "data.pdf"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="PDF"))
        assert res.success is True

    def test_xlsx_to_json(self, engine, tmp_path):
        src = _make_xlsx(tmp_path / "data.xlsx")
        dst = tmp_path / "data.json"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="JSON"))
        assert res.success is True

    def test_xlsx_to_html(self, engine, tmp_path):
        src = _make_xlsx(tmp_path / "data.xlsx")
        dst = tmp_path / "data.html"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="HTML"))
        assert res.success is True
        assert "<table" in dst.read_text(encoding="utf-8")

    def test_csv_metadata(self, tmp_path):
        src = _make_csv(tmp_path / "meta.csv")
        meta = get_csv_metadata(src)
        assert meta["format"] == "CSV"
        assert meta["rows"] == 2
        assert "Name" in meta["columns"]

    def test_xlsx_metadata(self, tmp_path):
        src = _make_xlsx(tmp_path / "meta.xlsx")
        meta = get_excel_metadata(src)
        assert meta["format"] == "XLSX"
        assert "Sheet1" in meta["sheets"]


# ============================================================================
# SECTION 6: Security Audit
# ============================================================================

class TestSecurity:

    def test_path_traversal_blocked(self, tmp_path):
        base = tmp_path / "sandbox"
        base.mkdir()
        bad_path = base / ".." / "escaped.png"
        with pytest.raises(SecurityError, match="Path traversal"):
            validate_output_path(bad_path, base_dir=base)

    def test_null_byte_blocked(self):
        with pytest.raises(SecurityError, match="Null byte"):
            validate_output_path("file\x00.png")

    def test_safe_output_path(self, tmp_path):
        safe = tmp_path / "sub" / "out.png"
        result = validate_output_path(safe, base_dir=tmp_path)
        assert result.resolve() == safe.resolve()

    def test_sniff_pdf(self, tmp_path):
        f = tmp_path / "t.pdf"
        f.write_bytes(b"%PDF-1.7\n%...")
        assert sniff_file_type(f) == "PDF"

    def test_sniff_png(self, tmp_path):
        f = tmp_path / "t.png"
        f.write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00")
        assert sniff_file_type(f) == "PNG"

    def test_sniff_jpeg(self, tmp_path):
        f = tmp_path / "t.jpg"
        f.write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF")
        assert sniff_file_type(f) == "JPEG"

    def test_sniff_rtf(self, tmp_path):
        f = tmp_path / "t.rtf"
        f.write_bytes(b"{\\rtf1\\ansi Hello}")
        assert sniff_file_type(f) == "RTF"

    def test_sniff_unknown(self, tmp_path):
        f = tmp_path / "t.xyz"
        f.write_bytes(b"random bytes")
        assert sniff_file_type(f) is None

    def test_sniff_empty_file(self, tmp_path):
        f = tmp_path / "empty.bin"
        f.touch()
        assert sniff_file_type(f) is None

    def test_zip_container_safe(self, tmp_path):
        z = tmp_path / "safe.zip"
        with zipfile.ZipFile(z, "w") as zf:
            zf.writestr("file.txt", "hello")
        assert validate_zip_container(z) is True

    def test_zip_bomb_too_many_entries(self, tmp_path):
        z = tmp_path / "many.zip"
        with zipfile.ZipFile(z, "w") as zf:
            for i in range(20):
                zf.writestr(f"f_{i}.txt", "A")
        with pytest.raises(SecurityError, match="contains 20 files"):
            validate_zip_container(z, max_entries=10)

    def test_zip_bomb_oversized(self, tmp_path):
        z = tmp_path / "big.zip"
        with zipfile.ZipFile(z, "w") as zf:
            zf.writestr("big.txt", "0" * 2000)
        with pytest.raises(SecurityError, match="uncompressed size exceeds limit"):
            validate_zip_container(z, max_uncompressed_bytes=500)

    def test_file_signature_mismatch(self, tmp_path):
        f = tmp_path / "fake.png"
        f.write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF....")  # JPEG bytes in .png
        with pytest.raises(ValueError, match="File signature mismatch"):
            validate_input_file(f)

    def test_corrupted_pdf_rejected(self, engine, tmp_path):
        f = tmp_path / "bad.pdf"
        f.write_bytes(b"NOT A PDF FILE AT ALL")
        out = tmp_path / "out.txt"
        res = engine.convert_single(f, out, ConversionConfig(target_format="TXT"))
        assert res.success is False
        assert "corrupted or invalid pdf" in res.error_message.lower()

    def test_empty_file_rejected(self, engine, tmp_path):
        f = tmp_path / "empty.pdf"
        f.touch()
        out = tmp_path / "out.txt"
        res = engine.convert_single(f, out, ConversionConfig(target_format="TXT"))
        assert res.success is False
        assert "empty (0 bytes)" in res.error_message.lower()

    def test_max_file_size_enforced(self, engine, tmp_path):
        f = tmp_path / "big.txt"
        f.write_bytes(b"X" * 2048)
        out = tmp_path / "out.pdf"
        cfg = ConversionConfig(target_format="PDF", max_file_size=1024)
        res = engine.convert_single(f, out, cfg)
        assert res.success is False
        assert "exceeds maximum allowed limit" in res.error_message.lower()


# ============================================================================
# SECTION 7: Edge Cases & Error Handling
# ============================================================================

class TestEdgeCases:

    def test_nonexistent_input(self, engine, tmp_path):
        res = engine.convert_single(tmp_path / "ghost.jpg", tmp_path / "out.png",
                                    ConversionConfig(target_format="PNG"))
        assert res.success is False
        assert "not found" in res.error_message.lower()

    def test_same_input_output_rejected(self, engine, tmp_path):
        f = _make_image(tmp_path / "same.jpg")
        res = engine.convert_single(f, f, ConversionConfig(target_format="JPG"))
        assert res.success is False
        assert "must be different" in res.error_message

    def test_unsupported_format(self, engine, tmp_path):
        src = _make_image(tmp_path / "src.png")
        dst = tmp_path / "out.xyz"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="XYZ_BAD"))
        assert res.success is False
        assert "unsupported" in res.error_message.lower() or "cannot convert" in res.error_message.lower()

    def test_encrypted_pdf_no_password(self, engine, tmp_path):
        pdf = tmp_path / "locked.pdf"
        doc = fitz.open()
        doc.new_page()
        doc.save(str(pdf), encryption=fitz.PDF_ENCRYPT_AES_256,
                 owner_pw="own", user_pw="usr123")
        doc.close()
        out = tmp_path / "out.txt"
        res = engine.convert_single(pdf, out, ConversionConfig(target_format="TXT"))
        assert res.success is False
        assert "encrypted" in res.error_message.lower() or "password" in res.error_message.lower()

    def test_encrypted_pdf_wrong_password(self, engine, tmp_path):
        pdf = tmp_path / "locked.pdf"
        doc = fitz.open()
        doc.new_page()
        doc.save(str(pdf), encryption=fitz.PDF_ENCRYPT_AES_256,
                 owner_pw="own", user_pw="correct")
        doc.close()
        out = tmp_path / "out.txt"
        res = engine.convert_single(pdf, out, ConversionConfig(
            target_format="TXT", password="wrong"))
        assert res.success is False
        assert "incorrect password" in res.error_message.lower()

    def test_encrypted_pdf_correct_password(self, engine, tmp_path):
        pdf = tmp_path / "locked.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Secret content")
        doc.save(str(pdf), encryption=fitz.PDF_ENCRYPT_AES_256,
                 owner_pw="own", user_pw="pass123")
        doc.close()
        out = tmp_path / "out.txt"
        res = engine.convert_single(pdf, out, ConversionConfig(
            target_format="TXT", password="pass123"))
        assert res.success is True
        assert "Secret content" in out.read_text(encoding="utf-8")

    def test_encrypted_docx_detection(self, tmp_path):
        f = tmp_path / "enc.docx"
        ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
        f.write_bytes(ole_header + b"Padding " * 20 + b"EncryptedPackage" + b" " * 100)
        with pytest.raises(ValueError, match="encrypted / password-protected"):
            check_docx_encryption(f)

    def test_non_utf8_csv_cp1252(self, engine, tmp_path):
        f = tmp_path / "latin.csv"
        raw = "Name,City\nRen\u00e9,Z\u00fcrich\n"
        f.write_bytes(raw.encode("cp1252"))
        out = tmp_path / "out.xlsx"
        res = engine.convert_single(f, out, ConversionConfig(target_format="XLSX"))
        assert res.success is True

    def test_non_utf8_csv_utf16(self, engine, tmp_path):
        f = tmp_path / "utf16.csv"
        raw = "Product,Qty\nWidget,100\n"
        f.write_bytes(raw.encode("utf-16"))
        out = tmp_path / "out.xlsx"
        res = engine.convert_single(f, out, ConversionConfig(target_format="XLSX"))
        assert res.success is True

    def test_unicode_content_in_pdf(self, engine, tmp_path):
        src = tmp_path / "uni.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Unicode test: cafe\u0301 na\u00efve re\u0301sume\u0301")
        doc.save(str(src))
        doc.close()
        dst = tmp_path / "uni.txt"
        res = engine.convert_single(src, dst, ConversionConfig(target_format="TXT"))
        assert res.success is True


# ============================================================================
# SECTION 8: Metadata Operations Audit
# ============================================================================

class TestMetadata:

    def test_image_exif_inspect_and_strip(self, tmp_path):
        src = _make_image_with_exif(tmp_path / "exif.jpg")
        meta = get_detailed_metadata(src)
        assert meta["has_metadata"] is True
        assert meta["fields_count"] >= 2

        clean = tmp_path / "clean.jpg"
        res = strip_file_metadata(src, clean)
        assert res["success"] is True
        meta2 = get_detailed_metadata(clean)
        assert meta2["has_metadata"] is False

    def test_pdf_metadata_strip(self, tmp_path):
        src = _make_pdf_with_meta(tmp_path / "meta.pdf")
        meta = get_detailed_metadata(src)
        assert meta["has_metadata"] is True
        assert meta["fields"]["Author"] == "AuditAuthor"

        clean = tmp_path / "clean.pdf"
        strip_pdf_metadata(src, clean)
        meta2 = get_detailed_metadata(clean)
        assert meta2["has_metadata"] is False

    def test_docx_metadata_strip(self, tmp_path):
        src = _make_docx_with_meta(tmp_path / "meta.docx")
        meta = get_detailed_metadata(src)
        assert meta["has_metadata"] is True

        clean = tmp_path / "clean.docx"
        strip_docx_metadata(src, clean)
        reopened = docx.Document(str(clean))
        assert reopened.core_properties.author in (None, "")

    def test_xlsx_metadata_strip(self, tmp_path):
        src = _make_xlsx_with_meta(tmp_path / "meta.xlsx")
        meta = get_detailed_metadata(src)
        assert meta["has_metadata"] is True

        clean = tmp_path / "clean.xlsx"
        strip_xlsx_metadata(src, clean)
        wb = openpyxl.load_workbook(str(clean))
        assert wb.properties.creator in (None, "")
        wb.close()

    def test_odt_metadata_strip(self, tmp_path):
        odt = tmp_path / "meta.odt"
        meta_xml = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<office:document-meta xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" '
            'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0">\n'
            '<office:meta>\n<dc:creator>OdtAuthor</dc:creator>\n</office:meta>\n'
            '</office:document-meta>'
        )
        with zipfile.ZipFile(str(odt), "w") as z:
            z.writestr("meta.xml", meta_xml)
            z.writestr("content.xml", "<office:document-content/>")

        meta = get_detailed_metadata(odt)
        assert meta["has_metadata"] is True
        assert meta["fields"]["Creator"] == "OdtAuthor"

        clean = tmp_path / "clean.odt"
        strip_odt_metadata(odt, clean)
        with zipfile.ZipFile(str(clean), "r") as z:
            assert "OdtAuthor" not in z.read("meta.xml").decode("utf-8")

    def test_engine_strip_metadata_method(self, tmp_path):
        engine = UniversalConverterEngine()
        src = _make_pdf_with_meta(tmp_path / "eng.pdf")
        clean = tmp_path / "eng_clean.pdf"
        res = engine.strip_metadata(src, clean)
        assert res.success is True

    def test_strip_metadata_preserves_content(self, tmp_path):
        src = _make_pdf_with_meta(tmp_path / "content.pdf")
        clean = tmp_path / "content_clean.pdf"
        strip_pdf_metadata(src, clean)
        d = fitz.open(str(clean))
        assert len(d) >= 1
        d.close()

    def test_strip_file_metadata_dispatch_image(self, tmp_path):
        src = _make_image_with_exif(tmp_path / "dispatch.jpg")
        clean = tmp_path / "dispatch_clean.jpg"
        res = strip_file_metadata(src, clean)
        assert res["success"] is True

    def test_get_detailed_metadata_structure_image(self, tmp_path):
        src = _make_image(tmp_path / "struct.png")
        meta = get_detailed_metadata(src)
        assert "has_metadata" in meta
        assert "fields" in meta
        assert "fields_count" in meta
        assert "warnings" in meta

    def test_get_detailed_metadata_structure_pdf(self, tmp_path):
        src = _make_pdf(tmp_path / "struct.pdf")
        meta = get_detailed_metadata(src)
        assert "has_metadata" in meta
        assert "fields" in meta


# ============================================================================
# SECTION 9: OCR Audit
# ============================================================================

class TestOCR:

    def test_ocr_available_returns_bool(self):
        assert isinstance(is_ocr_available(), bool)

    def test_ocr_status_message_nonempty(self):
        msg = get_ocr_status_message()
        assert isinstance(msg, str)
        assert len(msg) > 0

    def test_pdf_with_text_needs_no_ocr(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "This is a document with full selectable text content exceeding thirty chars.")
        assert pdf_needs_ocr(doc) is False
        doc.close()

    def test_pdf_image_only_needs_ocr(self):
        doc = fitz.open()
        doc.new_page()  # no text
        assert pdf_needs_ocr(doc) is True
        doc.close()

    def test_ocr_image_to_text_mocked(self, tmp_path):
        img = _make_image(tmp_path / "scan.png")
        with patch("pytesseract.image_to_string", return_value="OCR Result ABC"):
            text = ocr_image_to_text(img)
            assert text == "OCR Result ABC"

    def test_ocr_config_propagation(self):
        c = ConversionConfig(enable_ocr=True)
        assert c.enable_ocr is True
        c2 = ConversionConfig(enable_ocr=False)
        assert c2.enable_ocr is False


# ============================================================================
# SECTION 10: Batch Processing & Concurrency Audit
# ============================================================================

class TestBatchProcessing:

    def test_batch_empty(self, engine):
        results = engine.convert_batch([])
        assert results == []

    def test_batch_single(self, engine, tmp_path):
        src = _make_image(tmp_path / "s.png")
        dst = tmp_path / "d.jpg"
        tasks = [(src, dst, ConversionConfig(target_format="JPG"))]
        results = engine.convert_batch(tasks)
        assert len(results) == 1
        assert results[0].success is True

    def test_batch_multiple_workers(self, engine, tmp_path):
        tasks = []
        for i in range(6):
            s = _make_image(tmp_path / f"in_{i}.png")
            d = tmp_path / f"out_{i}.jpg"
            tasks.append((s, d, ConversionConfig(target_format="JPG")))

        results = engine.convert_batch(tasks, max_workers=3)
        assert len(results) == 6
        assert all(r.success for r in results)

    def test_batch_progress_callback(self, engine, tmp_path):
        tasks = []
        for i in range(4):
            s = _make_image(tmp_path / f"p_{i}.png")
            d = tmp_path / f"po_{i}.jpg"
            tasks.append((s, d, ConversionConfig(target_format="JPG")))

        events = []
        def on_progress(completed, total, res):
            events.append((completed, total, res.success))

        engine.convert_batch(tasks, progress_callback=on_progress, max_workers=2)
        assert len(events) == 4
        assert events[-1][0] == 4
        assert events[-1][1] == 4

    def test_batch_cancellation(self, engine, tmp_path):
        tasks = []
        for i in range(10):
            s = _make_image(tmp_path / f"c_{i}.png")
            d = tmp_path / f"co_{i}.jpg"
            tasks.append((s, d, ConversionConfig(target_format="JPG")))

        cancel_flag = False
        count = 0
        def on_progress(completed, total, res):
            nonlocal count, cancel_flag
            count += 1
            if count >= 2:
                cancel_flag = True

        results = engine.convert_batch(tasks, progress_callback=on_progress,
                                       cancel_check=lambda: cancel_flag, max_workers=1)
        assert len(results) < 10

    def test_batch_mixed_success_failure(self, engine, tmp_path):
        s_good = _make_image(tmp_path / "good.png")
        d_good = tmp_path / "good.jpg"

        s_bad = tmp_path / "bad.pdf"
        s_bad.write_bytes(b"CORRUPT")
        d_bad = tmp_path / "bad.txt"

        tasks = [
            (s_good, d_good, ConversionConfig(target_format="JPG")),
            (s_bad, d_bad, ConversionConfig(target_format="TXT")),
        ]
        results = engine.convert_batch(tasks, max_workers=1)
        assert len(results) == 2
        successes = [r for r in results if r.success]
        failures = [r for r in results if not r.success]
        assert len(successes) >= 1
        assert len(failures) >= 1


# ============================================================================
# SECTION 11: CLI Audit
# ============================================================================

class TestCLI:

    def test_parse_color_hex6(self):
        assert parse_color("#ff8000") == (255, 128, 0)

    def test_parse_color_hex3(self):
        assert parse_color("#f00") == (255, 0, 0)

    def test_parse_color_rgb(self):
        assert parse_color("100, 200, 50") == (100, 200, 50)

    def test_parse_color_invalid(self):
        with pytest.raises(ValueError):
            parse_color("notacolor")

    def test_parse_color_out_of_range(self):
        with pytest.raises(ValueError):
            parse_color("256,0,0")

    def test_parse_color_too_many_channels(self):
        with pytest.raises(ValueError):
            parse_color("1,2,3,4")

    def test_collect_files_single(self, tmp_path):
        f = tmp_path / "a.png"
        _make_image(f)
        found = collect_convertible_files(f)
        assert len(found) == 1

    def test_collect_files_directory(self, tmp_path):
        _make_image(tmp_path / "a.png")
        _make_image(tmp_path / "b.jpg")
        found = collect_convertible_files(tmp_path)
        assert len(found) >= 2

    def test_collect_files_empty_dir(self, tmp_path):
        empty = tmp_path / "empty"
        empty.mkdir()
        found = collect_convertible_files(empty)
        assert len(found) == 0

    def test_cli_list_formats(self, capsys):
        code = main(["--list-formats"])
        assert code == 0
        out = capsys.readouterr().out
        assert "Images" in out
        assert "Documents" in out
        assert "Spreadsheets" in out

    def test_cli_single_conversion(self, tmp_path):
        src = _make_image(tmp_path / "cli_in.png")
        dst = tmp_path / "cli_out.jpg"
        code = main(["-i", str(src), "-o", str(dst)])
        assert code == 0
        assert dst.is_file()

    def test_cli_pdf_to_txt(self, tmp_path):
        src = _make_pdf(tmp_path / "cli.pdf", "CLI PDF text")
        dst = tmp_path / "cli.txt"
        code = main(["-i", str(src), "-o", str(dst)])
        assert code == 0
        assert "CLI PDF text" in dst.read_text(encoding="utf-8")

    def test_cli_inspect_metadata(self, tmp_path, capsys):
        src = _make_pdf_with_meta(tmp_path / "meta.pdf")
        code = main(["--inspect-metadata", str(src)])
        assert code == 0
        out = capsys.readouterr().out
        assert "Metadata Inspection" in out
        assert "AuditAuthor" in out

    def test_cli_clean_metadata(self, tmp_path):
        src = _make_pdf_with_meta(tmp_path / "to_clean.pdf")
        code = main(["--clean-metadata", str(src)])
        assert code == 0
        clean = tmp_path / "to_clean_clean.pdf"
        assert clean.exists()
        meta = get_detailed_metadata(clean)
        assert meta["has_metadata"] is False

    def test_cli_strip_metadata_during_convert(self, tmp_path):
        src = _make_image_with_exif(tmp_path / "strip.jpg")
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        code = main(["-i", str(src), "-f", "PNG", "-o", str(out_dir), "--strip-metadata"])
        assert code == 0
        out_file = out_dir / "strip.png"
        assert out_file.exists()
        meta = get_detailed_metadata(out_file)
        assert meta["has_metadata"] is False

    def test_cli_rejects_missing_input(self):
        code = main(["-i", "nonexistent_file_xyz.jpg", "-f", "PNG"])
        assert code == 1

    def test_cli_rejects_invalid_format(self, tmp_path):
        src = _make_image(tmp_path / "fmt.png")
        code = main(["-i", str(src), "-f", "BADFORMAT"])
        assert code == 1

    def test_cli_rejects_duplicate_destinations(self, tmp_path, capsys):
        d_in = tmp_path / "inputs"
        d_out = tmp_path / "outputs"
        d_in.mkdir()
        d_out.mkdir()
        _make_image(d_in / "same.jpg")
        _make_image(d_in / "same.png")
        code = main(["-i", str(d_in), "-f", "JPG", "-o", str(d_out)])
        assert code == 1
        assert "would overwrite" in capsys.readouterr().err

    def test_cli_no_input_shows_help(self, capsys):
        code = main([])
        assert code == 1


# ============================================================================
# SECTION 12: GUI Widget Audit
# ============================================================================

HAS_PYSIDE = False
try:
    from PySide6.QtWidgets import QApplication
    HAS_PYSIDE = True
except ImportError:
    pass

_qapp = None

def get_qapp():
    global _qapp
    if _qapp is None:
        _qapp = QApplication.instance()
        if _qapp is None:
            _qapp = QApplication([])
    return _qapp


@pytest.mark.skipif(not HAS_PYSIDE, reason="PySide6 not installed")
class TestGUIWidgets:

    @pytest.fixture(autouse=True)
    def _ensure_qapp(self):
        get_qapp()

    def test_nav_rail_has_3_buttons(self):
        from image_converter.ui.nav_rail import NavRailWidget
        widget = NavRailWidget()
        assert len(widget.buttons) == 3

    def test_nav_rail_tab_changed_signal(self):
        from image_converter.ui.nav_rail import NavRailWidget
        widget = NavRailWidget()
        received = []
        widget.tab_changed.connect(lambda idx: received.append(idx))
        widget.set_active_tab(1)
        assert 1 in received

    def test_nav_rail_set_active_tab(self):
        from image_converter.ui.nav_rail import NavRailWidget
        widget = NavRailWidget()
        widget.set_active_tab(2)
        assert widget.buttons[2].isChecked()

    def test_stats_card_set_value(self):
        from image_converter.ui.stats_widget import StatsCard
        card = StatsCard("Test", "0", "📈")
        card.set_value("42", "#ff0000")
        assert card.val_lbl.text() == "42"

    def test_stats_widget_update(self):
        from image_converter.ui.stats_widget import StatsWidget
        sw = StatsWidget()
        sw.update_stats(5, 1, 2.5, 1024 * 100, 1024 * 80)
        assert sw.card_succeeded.val_lbl.text() == "5"
        assert sw.card_failed.val_lbl.text() == "1"
        assert "2.50s" in sw.card_time.val_lbl.text()

    def test_format_bytes_utility(self):
        from image_converter.ui.components import format_bytes
        assert format_bytes(0) == "0 B"
        assert "B" in format_bytes(500)
        assert "KB" in format_bytes(1024)
        assert "MB" in format_bytes(1024 * 1024)
        assert "MB" in format_bytes(1024 ** 3)

    def test_conversion_settings_get_config(self):
        from image_converter.ui.components import ConversionSettingsWidget
        widget = ConversionSettingsWidget()
        cfg = widget.get_config()
        assert isinstance(cfg, ConversionConfig)

    def test_drop_zone_instantiation(self):
        from image_converter.ui.components import DropZoneWidget
        widget = DropZoneWidget()
        assert widget is not None

    def test_queue_table_instantiation(self):
        from image_converter.ui.components import QueueTableWidget
        widget = QueueTableWidget()
        assert widget is not None

    def test_queue_table_per_row_target_format(self, tmp_path):
        from image_converter.ui.components import QueueTableWidget
        widget = QueueTableWidget()
        test_file = tmp_path / "sample.docx"
        test_file.write_text("dummy")
        widget.add_file_item(test_file, "PDF")
        assert widget.rowCount() == 1
        assert widget.get_target_format_for_row(0) == "PDF"
        widget.set_target_format_for_row(0, "TXT")
        assert widget.get_target_format_for_row(0) == "TXT"

    def test_main_window_reentrancy_lock(self):
        from image_converter.ui.main_window import ImageConverterMainWindow
        win = ImageConverterMainWindow()
        assert win._is_converting is False
        win._is_converting = True
        # start_conversion should return early without doing anything
        win.start_conversion()
        assert win._is_converting is True

    def test_main_window_instantiation(self):
        from image_converter.ui.main_window import ImageConverterMainWindow
        win = ImageConverterMainWindow()
        assert hasattr(win, "nav_rail")
        assert hasattr(win, "stats_widget")
        assert hasattr(win, "drop_zone")
        assert hasattr(win, "queue_table")


# ============================================================================
# SECTION 13: Logging Audit
# ============================================================================

class TestLogging:

    def test_setup_logger_creates_logger(self, tmp_path):
        log_file = tmp_path / "test.log"
        logger = setup_logger(verbose=True, log_file=log_file)
        assert logger is not None

    def test_get_logger_namespaced(self):
        lg = get_logger("audit")
        assert lg.name == "universal_converter.audit"

    def test_logger_writes_to_file(self, tmp_path):
        log_file = tmp_path / "write.log"
        logger = setup_logger(verbose=True, log_file=log_file)
        logger.info("QA_AUDIT_LOG_MSG_12345")
        for h in logger.handlers:
            h.flush()
        assert log_file.is_file()
        assert "QA_AUDIT_LOG_MSG_12345" in log_file.read_text(encoding="utf-8")


# ============================================================================
# SECTION 14: Cross-Cutting Integration
# ============================================================================

class TestIntegration:

    def test_engine_has_all_converters(self):
        eng = UniversalConverterEngine()
        assert hasattr(eng, "doc_converter")
        assert hasattr(eng, "data_converter")
        assert hasattr(eng, "pres_converter")
        assert hasattr(eng, "rich_converter")
        assert isinstance(eng.doc_converter, DocumentConverter)
        assert isinstance(eng.data_converter, DataConverter)
        assert isinstance(eng.pres_converter, PresentationConverter)
        assert isinstance(eng.rich_converter, RichDocumentConverter)

    def test_full_pipeline_create_convert_strip(self, tmp_path):
        """Full pipeline: create -> convert -> inspect -> strip -> verify."""
        engine = UniversalConverterEngine()
        # Create image with EXIF
        src = _make_image_with_exif(tmp_path / "pipeline.jpg")

        # Convert to PNG
        png = tmp_path / "pipeline.png"
        r1 = engine.convert_single(src, png, ConversionConfig(target_format="PNG"))
        assert r1.success is True

        # Inspect metadata on converted
        meta = get_detailed_metadata(png)
        assert isinstance(meta, dict)

        # Strip metadata
        clean = tmp_path / "pipeline_clean.png"
        r2 = strip_file_metadata(png, clean)
        assert r2["success"] is True

        # Verify clean
        meta2 = get_detailed_metadata(clean)
        assert meta2["has_metadata"] is False

    def test_batch_with_strip_metadata(self, tmp_path):
        engine = UniversalConverterEngine()
        tasks = []
        for i in range(3):
            s = _make_image_with_exif(tmp_path / f"b_{i}.jpg")
            d = tmp_path / f"b_{i}.png"
            tasks.append((s, d, ConversionConfig(target_format="PNG", strip_metadata=True)))
        results = engine.convert_batch(tasks, max_workers=2)
        assert len(results) == 3
        assert all(r.success for r in results)

    def test_resize_quality_format_simultaneous(self, tmp_path):
        """Convert with resize + quality + format change simultaneously."""
        engine = UniversalConverterEngine()
        src = _make_image(tmp_path / "multi.png", w=300, h=200)
        dst = tmp_path / "multi.jpg"
        cfg = ConversionConfig(
            target_format="JPG", quality=50,
            resize_mode="percentage", resize_percent=50.0)
        res = engine.convert_single(src, dst, cfg)
        assert res.success is True
        assert res.output_dimensions == (150, 100)

    def test_round_trip_preserves_content(self, tmp_path):
        """Round trip: PDF -> TXT -> PDF keeps text."""
        engine = UniversalConverterEngine()
        src = _make_pdf(tmp_path / "rt.pdf", "Round trip content")
        txt = tmp_path / "rt.txt"
        res1 = engine.convert_single(src, txt, ConversionConfig(target_format="TXT"))
        assert res1.success is True
        assert "Round trip content" in txt.read_text(encoding="utf-8")

        back = tmp_path / "rt_back.pdf"
        res2 = engine.convert_single(txt, back, ConversionConfig(target_format="PDF"))
        assert res2.success is True
        assert back.stat().st_size > 0

    def test_supported_formats_consistency(self):
        """Input extensions and output formats are consistent."""
        in_exts = get_supported_input_extensions()
        out_fmts = get_supported_output_formats()
        assert len(in_exts) > 10
        assert len(out_fmts) > 10
        assert ".jpg" in in_exts
        assert ".pdf" in in_exts
        assert ".csv" in in_exts
        assert "PNG" in out_fmts
        assert "PDF" in out_fmts
        assert "XLSX" in out_fmts
