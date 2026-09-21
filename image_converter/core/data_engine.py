"""
Data conversion engine for CSV and Excel (.xlsx, .xls) files.
Supports converting between CSV, Excel, PDF tables, JSON, HTML, and Text.
"""

from __future__ import annotations

import csv
import html
import json
import logging
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd

logger = logging.getLogger(__name__)

# Check openpyxl
_OPENPYXL_AVAILABLE = False
try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    _OPENPYXL_AVAILABLE = True
except ImportError:
    pass

# Check reportlab
_REPORTLAB_AVAILABLE = False
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table as RLTable,
        TableStyle as RLTableStyle,
    )
    _REPORTLAB_AVAILABLE = True
except ImportError:
    pass


def read_csv_with_fallback(file_path: str | Path, **kwargs) -> pd.DataFrame:
    """
    Read a CSV file with automatic encoding detection and fallback list
    to reliably handle UTF-8, UTF-8-SIG, CP1252, Latin-1, ISO-8859-1, and UTF-16.
    """
    p = Path(file_path).resolve()
    encodings_to_try: List[str] = []

    # 1. Check BOM signatures first
    try:
        with open(p, "rb") as f:
            bom = f.read(4)
        if bom.startswith(b"\xff\xfe") or bom.startswith(b"\xfe\xff"):
            encodings_to_try.append("utf-16")
        elif bom.startswith(b"\xef\xbb\xbf"):
            encodings_to_try.append("utf-8-sig")
    except Exception as e:
        logger.warning(f"Error reading BOM signatures for CSV {p.name}: {e}", exc_info=True)

    # 2. Standard common encodings
    encodings_to_try.extend(["utf-8", "cp1252", "latin-1", "iso-8859-1"])

    # 3. Charset normalizer as auxiliary detector
    try:
        from charset_normalizer import from_path
        results = from_path(p)
        best = results.best()
        if best and best.encoding:
            encodings_to_try.append(best.encoding)
    except Exception as e:
        logger.warning(f"Charset detection error for CSV {p.name}: {e}", exc_info=True)

    encodings_to_try.append("utf-16")

    # Deduplicate while preserving order
    seen = set()
    unique_encodings = [e for e in encodings_to_try if not (e in seen or seen.add(e))]

    for enc in unique_encodings:
        try:
            df = pd.read_csv(str(p), encoding=enc, **kwargs)
            if any("\x00" in str(col) for col in df.columns):
                continue
            return df
        except (UnicodeDecodeError, UnicodeError) as enc_err:
            logger.debug(f"Candidate encoding '{enc}' failed for CSV '{p.name}': {enc_err}")
            continue

    # Final fallback replacing invalid characters
    return pd.read_csv(str(p), encoding="utf-8", encoding_errors="replace", **kwargs)


def get_csv_metadata(file_path: str | Path) -> Dict[str, Any]:
    """Extract metadata, row count, and column names from a CSV file."""
    p = Path(file_path).resolve()
    if not p.is_file():
        raise FileNotFoundError(f"CSV file not found: {file_path}")

    file_size = p.stat().st_size
    rows = 0
    cols = 0
    col_names: List[str] = []

    try:
        df = read_csv_with_fallback(p, nrows=100)
        cols = len(df.columns)
        col_names = [str(c) for c in df.columns]

        # Fast line count
        with open(p, "rb") as f:
            rows = max(0, sum(1 for _ in f) - 1)
    except Exception as e:
        logger.warning(f"Error extracting CSV metadata for {p.name}: {e}", exc_info=True)

    return {
        "file_name": p.name,
        "file_path": str(p),
        "file_size": file_size,
        "width": cols,
        "height": rows,
        "rows": rows,
        "cols": cols,
        "columns": col_names[:10],
        "format": "CSV",
        "has_alpha": False,
    }


def get_excel_metadata(file_path: str | Path) -> Dict[str, Any]:
    """Extract metadata, sheet names, and dimensions from an Excel file."""
    p = Path(file_path).resolve()
    if not p.is_file():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    file_size = p.stat().st_size
    sheet_names: List[str] = []
    total_rows = 0
    total_cols = 0

    if _OPENPYXL_AVAILABLE:
        try:
            wb = openpyxl.load_workbook(str(p), read_only=True)
            sheet_names = wb.sheetnames
            if sheet_names:
                ws = wb[sheet_names[0]]
                total_rows = ws.max_row or 0
                total_cols = ws.max_column or 0
            wb.close()
        except Exception as e:
            logger.warning(f"openpyxl failed to read Excel metadata for {p.name}: {e}", exc_info=True)

    if not sheet_names:
        try:
            xl = pd.ExcelFile(str(p))
            sheet_names = xl.sheet_names
            df = xl.parse(sheet_names[0], nrows=5)
            total_rows = len(df)
            total_cols = len(df.columns)
        except Exception as e:
            logger.warning(f"pandas fallback failed to read Excel metadata for {p.name}: {e}", exc_info=True)

    return {
        "file_name": p.name,
        "file_path": str(p),
        "file_size": file_size,
        "width": total_cols,
        "height": total_rows,
        "rows": total_rows,
        "cols": total_cols,
        "sheets": sheet_names,
        "format": "XLSX",
        "has_alpha": False,
    }


class DataConverter:
    """Core data conversion engine handling CSV and Excel spreadsheets."""

    def __init__(self):
        self.has_openpyxl = _OPENPYXL_AVAILABLE
        self.has_reportlab = _REPORTLAB_AVAILABLE

    # -------------------------------------------------------------------------
    # CSV Input Conversions
    # -------------------------------------------------------------------------

    def convert_csv_to_excel(
        self,
        csv_path: str | Path,
        output_path: str | Path,
        sheet_name: str = "Data",
    ) -> Path:
        """
        Convert CSV to Excel (.xlsx) with auto-formatted column widths and header styling.
        """
        in_p = Path(csv_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        df = read_csv_with_fallback(in_p)

        if self.has_openpyxl:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name

            # Header styling
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_align = Alignment(horizontal="center", vertical="center")

            # Write header
            headers = [str(c) for c in df.columns]
            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align

            # Write data rows
            for row in df.itertuples(index=False):
                ws.append(list(row))

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(10, min(max_len + 3, 50))

            wb.save(str(out_p))
        else:
            df.to_excel(str(out_p), sheet_name=sheet_name, index=False)

        return out_p

    def convert_csv_to_pdf(
        self,
        csv_path: str | Path,
        output_path: str | Path,
        title: Optional[str] = None,
        max_rows: Optional[int] = 1000,
    ) -> Path:
        """
        Convert CSV data into a cleanly styled printable PDF table.
        Automatically uses landscape layout when the table has more than 5 columns.
        """
        if not self.has_reportlab:
            raise RuntimeError("reportlab is required for PDF table generation.")

        in_p = Path(csv_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        df = read_csv_with_fallback(in_p)
        if max_rows and len(df) > max_rows:
            df = df.iloc[:max_rows]

        return self._dataframe_to_pdf(df, out_p, title or in_p.stem)

    def convert_csv_to_json(
        self,
        csv_path: str | Path,
        output_path: str | Path,
        indent: int = 2,
    ) -> Path:
        """Convert CSV to JSON records."""
        in_p = Path(csv_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        df = read_csv_with_fallback(in_p)
        df.to_json(str(out_p), orient="records", indent=indent, date_format="iso")
        return out_p

    def convert_csv_to_html(
        self,
        csv_path: str | Path,
        output_path: str | Path,
    ) -> Path:
        """Convert CSV to a standalone responsive HTML page with modern styling."""
        in_p = Path(csv_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        df = read_csv_with_fallback(in_p)
        return self._dataframe_to_html(df, out_p, in_p.stem)

    def convert_csv_to_text(
        self,
        csv_path: str | Path,
        output_path: str | Path,
        delimiter: str = "\t",
    ) -> Path:
        """Convert CSV to tab-separated or formatted text file."""
        in_p = Path(csv_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        df = read_csv_with_fallback(in_p)
        df.to_csv(str(out_p), sep=delimiter, index=False)
        return out_p

    # -------------------------------------------------------------------------
    # Excel (XLSX) Input Conversions
    # -------------------------------------------------------------------------

    def convert_excel_to_csv(
        self,
        xlsx_path: str | Path,
        output_path: str | Path,
        sheet_name: Optional[str] = None,
    ) -> Path:
        """
        Convert Excel spreadsheet to CSV. If sheet_name is None, converts active/first sheet.
        """
        in_p = Path(xlsx_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        # Read sheet
        target_sheet = sheet_name or 0
        df = pd.read_excel(str(in_p), sheet_name=target_sheet)
        df.to_csv(str(out_p), index=False, encoding="utf-8-sig")
        return out_p

    def convert_excel_to_pdf(
        self,
        xlsx_path: str | Path,
        output_path: str | Path,
        sheet_name: Optional[str] = None,
    ) -> Path:
        """Convert Excel sheet to printable PDF report natively via COM."""
        in_p = Path(xlsx_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        if os.name == "nt":
            import time
            from image_converter.core.com_utils import (
                com_initialized,
                EXCEL_ALERTS_NONE,
                MSO_AUTOMATION_SECURITY_FORCE_DISABLE,
                XL_TYPE_PDF,
            )
            with com_initialized():
                import win32com.client
                try:
                    excel = win32com.client.DispatchEx("Excel.Application")
                except Exception as dispatch_err:
                    logger.warning(f"Excel COM not available: {dispatch_err}")
                    excel = None

                if excel is not None:
                    try:
                        excel.Visible = False
                        excel.DisplayAlerts = EXCEL_ALERTS_NONE
                        excel.AutomationSecurity = MSO_AUTOMATION_SECURITY_FORCE_DISABLE
                        wb = None
                        tmp_out = out_p.with_name(f"{out_p.stem}_tmp_{os.getpid()}_{time.time_ns()}.pdf")
                        try:
                            try:
                                wb = excel.Workbooks.Open(str(in_p), ReadOnly=True)
                                if sheet_name:
                                    ws = wb.Sheets(sheet_name)
                                    ws.ExportAsFixedFormat(XL_TYPE_PDF, str(tmp_out))
                                else:
                                    wb.ExportAsFixedFormat(XL_TYPE_PDF, str(tmp_out))
                            finally:
                                if wb is not None:
                                    try:
                                        wb.Close(False)
                                    except Exception as close_err:
                                        logger.warning(f"Error closing Excel workbook: {close_err}")
                            if tmp_out.is_file() and tmp_out.stat().st_size > 0:
                                os.replace(tmp_out, out_p)
                        finally:
                            if tmp_out.exists():
                                try:
                                    tmp_out.unlink()
                                except Exception as unlink_err:
                                    logger.warning(f"Could not remove temporary file {tmp_out}: {unlink_err}")
                    except Exception as e:
                        logger.warning(f"MS Excel COM conversion failed for {in_p.name}: {e}", exc_info=True)
                        raise RuntimeError(f"Microsoft Excel conversion failed for '{in_p.name}': {e}") from e
                    finally:
                        ws = None
                        wb = None
                        try:
                            excel.Quit()
                        except Exception as quit_err:
                            logger.warning(f"Error quitting Excel application: {quit_err}")
                        excel = None
                        import gc
                        gc.collect()
                    if out_p.is_file() and out_p.stat().st_size > 0:
                        return out_p

        raise RuntimeError("Native Excel to PDF conversion requires Microsoft Excel on Windows.")

    def convert_excel_to_json(
        self,
        xlsx_path: str | Path,
        output_path: str | Path,
        sheet_name: Optional[str] = None,
    ) -> Path:
        """Convert Excel sheet or all sheets to JSON."""
        in_p = Path(xlsx_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        if sheet_name is not None:
            df = pd.read_excel(str(in_p), sheet_name=sheet_name)
            df.to_json(str(out_p), orient="records", indent=2, date_format="iso")
        else:
            # Read all sheets into dictionary
            sheets_dict = pd.read_excel(str(in_p), sheet_name=None)
            json_dict = {}
            for name, sheet_df in sheets_dict.items():
                json_dict[name] = json.loads(sheet_df.to_json(orient="records", date_format="iso"))
            out_p.write_text(json.dumps(json_dict, indent=2), encoding="utf-8")

        return out_p

    def convert_excel_to_html(
        self,
        xlsx_path: str | Path,
        output_path: str | Path,
        sheet_name: Optional[str] = None,
    ) -> Path:
        """Convert Excel sheet(s) to modern responsive HTML tables."""
        in_p = Path(xlsx_path).resolve()
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)

        if sheet_name is not None:
            df = pd.read_excel(str(in_p), sheet_name=sheet_name)
            return self._dataframe_to_html(df, out_p, f"{in_p.stem} - {sheet_name}")
        else:
            sheets_dict = pd.read_excel(str(in_p), sheet_name=None)
            html_parts = []
            for s_name, s_df in sheets_dict.items():
                html_parts.append(f"<h2>Sheet: {html.escape(s_name)}</h2>")
                tbl_html = s_df.to_html(classes="data-table", index=False, border=0)
                html_parts.append(tbl_html)

            full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>{html.escape(in_p.stem)}</title>
  <style>
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      margin: 30px;
      color: #1f2937;
    }}
    h1, h2 {{ color: #111827; }}
    .data-table {{
      width: 100%;
      border-collapse: collapse;
      margin: 16px 0 32px 0;
      font-size: 14px;
    }}
    .data-table th {{
      background-color: #1e293b;
      color: white;
      text-align: left;
      padding: 10px 12px;
    }}
    .data-table td {{
      padding: 8px 12px;
      border-bottom: 1px solid #e5e7eb;
    }}
    .data-table tr:nth-child(even) {{
      background-color: #f9fafb;
    }}
  </style>
</head>
<body>
  <h1>{html.escape(in_p.stem)}</h1>
  {chr(10).join(html_parts)}
</body>
</html>"""
            out_p.write_text(full_html, encoding="utf-8")
            return out_p

    # -------------------------------------------------------------------------
    # Shared Helpers
    # -------------------------------------------------------------------------

    def _dataframe_to_pdf(self, df: pd.DataFrame, out_path: Path, title: str) -> Path:
        """Render DataFrame into styled ReportLab PDF."""
        styles = getSampleStyleSheet()
        num_cols = len(df.columns)
        page_size = landscape(letter) if num_cols > 5 else letter

        # Build table data
        cell_style = ParagraphStyle(
            "CellText",
            parent=styles["Normal"],
            fontSize=8 if num_cols > 8 else 9,
            leading=10 if num_cols > 8 else 11,
            textColor=colors.HexColor("#1e293b"),
        )
        head_style = ParagraphStyle(
            "HeadText",
            parent=styles["Normal"],
            fontSize=8 if num_cols > 8 else 9,
            leading=10 if num_cols > 8 else 11,
            textColor=colors.white,
            fontName="Helvetica-Bold",
        )

        table_data = []
        # Header row
        table_data.append([Paragraph(html.escape(str(col)), head_style) for col in df.columns])

        # Rows
        for _, row in df.iterrows():
            row_cells = []
            for val in row:
                val_str = "" if pd.isna(val) else str(val)
                row_cells.append(Paragraph(html.escape(val_str), cell_style))
            table_data.append(row_cells)

        rl_table = RLTable(table_data, repeatRows=1)
        rl_table.setStyle(RLTableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=10,
        )

        story = [
            Paragraph(html.escape(title), title_style),
            Spacer(1, 8),
            rl_table,
        ]

        doc = SimpleDocTemplate(
            str(out_path),
            pagesize=page_size,
            leftMargin=30,
            rightMargin=30,
            topMargin=30,
            bottomMargin=30,
        )
        doc.build(story)
        return out_path

    def _dataframe_to_html(self, df: pd.DataFrame, out_path: Path, title: str) -> Path:
        """Render DataFrame into standalone HTML."""
        tbl_html = df.to_html(classes="data-table", index=False, border=0)
        full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>{html.escape(title)}</title>
  <style>
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      margin: 40px auto;
      max-width: 1200px;
      padding: 0 20px;
      color: #1f2937;
    }}
    h1 {{ color: #111827; }}
    .data-table {{
      width: 100%;
      border-collapse: collapse;
      margin: 20px 0;
      font-size: 14px;
      box-shadow: 0 1px 3px 0 rgba(0, 0, 0, 0.1);
      border-radius: 8px;
      overflow: hidden;
    }}
    .data-table th {{
      background-color: #1e293b;
      color: white;
      text-align: left;
      padding: 12px 14px;
      font-weight: 600;
    }}
    .data-table td {{
      padding: 10px 14px;
      border-bottom: 1px solid #f1f5f9;
    }}
    .data-table tr:nth-child(even) {{
      background-color: #f8fafc;
    }}
    .data-table tr:hover {{
      background-color: #f1f5f9;
    }}
  </style>
</head>
<body>
  <h1>{html.escape(title)}</h1>
  {tbl_html}
</body>
</html>"""
        out_path.write_text(full_html, encoding="utf-8")
        return out_path
